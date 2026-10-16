"""Constrói a folha de medições N2408 a partir dos JSON extraídos (scripts/extrair.py).

Cadeia de rastreabilidade: Resumo <- folhas por família (SUMIFS) <- Rastreio (1 linha por
entidade CAD, com handle) <- DXF <- DWG no Google Drive (ID no separador Ficheiros).
"""
import csv, glob, json, os, re, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule

RAIZ = '/home/user/Pedro/N2408'
SAIDA = f'{RAIZ}/N2408_Medicoes_DWG.xlsx'

# ---------- mapeamento desenho -> família e piso ----------
PISO = {
    'ARQ/01-Desenhos Gerais/1-01': 'Piso 0', 'ARQ/01-Desenhos Gerais/1-02': 'Piso 1', 'ARQ/01-Desenhos Gerais/1-03': 'Piso 2',
    'ARQ/01-Desenhos Gerais/1-04': 'Piso 3', 'ARQ/01-Desenhos Gerais/1-05': 'Cobertura', 'ARQ/01-Desenhos Gerais/1-06': 'Ruína (P1/P2)',
    'ARQ/02-Infraestruturas/2-06': 'Piso 0', 'ARQ/02-Infraestruturas/2-07': 'Piso 1', 'ARQ/02-Infraestruturas/2-08': 'Piso 2',
    'ARQ/02-Infraestruturas/2-09': 'Piso 3', 'ARQ/02-Infraestruturas/2-10': 'Ruína (P1/P2)', 'ARQ/02-Infraestruturas/2-11': 'Exterior',
    'ARQ/03-Acabamentos/Pavimentos/3-01': 'Piso 0', 'ARQ/03-Acabamentos/Pavimentos/3-02': 'Piso 1',
    'ARQ/03-Acabamentos/Pavimentos/3-03': 'Piso 2', 'ARQ/03-Acabamentos/Pavimentos/3-04': 'Piso 3',
    'ARQ/03-Acabamentos/Paredes/3-05': 'Piso 0', 'ARQ/03-Acabamentos/Paredes/3-06': 'Piso 1', 'ARQ/03-Acabamentos/Paredes/3-07': 'Piso 2',
    'ARQ/03-Acabamentos/Paredes/3-08': 'Piso 3', 'ARQ/03-Acabamentos/Paredes/3-09': 'Ruína (P1/P2)',
    'ARQ/03-Acabamentos/Tetos/3-10': 'Piso 0', 'ARQ/03-Acabamentos/Tetos/3-11': 'Piso 1', 'ARQ/03-Acabamentos/Tetos/3-12': 'Piso 2',
    'ARQ/03-Acabamentos/Tetos/3-13': 'Piso 3',
    'ARQ/05-Vãos Exteriores/5-05': 'Piso 1', 'ARQ/05-Vãos Exteriores/5-06': 'Piso 2',
}
AVAC_PISO = {'Piso -1': 'Piso 0 (AVAC: Piso -1)', 'Piso 1': 'Piso 1', 'Piso 2': 'Piso 2', 'Piso 3': 'Piso 3', 'Cobertura': 'Cobertura', 'Coberturas': 'Cobertura'}
PISOS = ['Piso 0', 'Piso 1', 'Piso 2', 'Piso 3', 'Cobertura', 'Ruína (P1/P2)', 'Exterior']


def familia(rel):
    if '/Pavimentos/' in rel: return 'PAV'
    if '/Tetos/' in rel: return 'TET'
    if '/Paredes/' in rel: return 'PAR'
    if '02-Infraestruturas' in rel: return 'INF'
    if rel.startswith('AVAC/'):
        return 'AVAC' if ('Aeraulicas' in rel or 'Tubagem' in rel) else 'AVAC-ESQ'
    return 'OUT'


def piso_de(rel):
    if rel in PISO:
        return PISO[rel]
    m = re.search(r'Planta d[oe]s? (Piso -?\d|Coberturas?)', rel)
    if m:
        return AVAC_PISO.get(m.group(1), m.group(1))
    return '-'


EXTERIOR = re.compile(r'DECK|CUBO DE GRANITO|TERRENO VEGETAL|PROTEC|CADEIRAS LOUNGE|DEGRAUS EM CHAPA|PASTILHA', re.I)

# ---------- estilos ----------
F = 'Arial'
fnt = Font(name=F, size=10)
fb = Font(name=F, size=10, bold=True)
fh = Font(name=F, size=10, bold=True, color='FFFFFF')
ft = Font(name=F, size=14, bold=True, color='1F3864')
fin = Font(name=F, size=10, color='0000FF')        # dados extraídos do CAD (entrada)
fln = Font(name=F, size=10, color='008000')        # ligação a outra folha
fill_h = PatternFill('solid', fgColor='1F3864')
fill_in = PatternFill('solid', fgColor='FFFF00')
fill_sub = PatternFill('solid', fgColor='D9E1F2')
fill_warn = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='BFBFBF')
brd = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')


def cab(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font, x.fill, x.alignment, x.border = fh, fill_h, Alignment(wrap_text=True, vertical='center'), brd
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def tabela(ws, ref, nome):
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    ws.add_table(t)


# ---------- ler dados ----------
man = {}
for r in csv.DictReader(open(f'{RAIZ}/manifesto.tsv'), delimiter='\t'):
    rel = r['pasta'] + '/' + r['ficheiro'][:-4]
    man[rel] = r

dados = {}
for jf in sorted(glob.glob(f'{RAIZ}/json/**/*.json', recursive=True)):
    rel = jf[len(RAIZ) + 6:-5]
    try:
        dados[rel] = json.load(open(jf))
    except Exception as e:
        print('ERRO', jf, e)

def tipo_avac(rel):
    return 'TUB' if 'Tubagem' in rel else ('AER' if 'Aeraulicas' in rel else 'OUT')


AVAC_COR = collections.defaultdict(set)   # (tipo de desenho, cor) -> descrições
for rel, d in dados.items():
    if rel.startswith('AVAC/'):
        for k, v in d['legenda'].items():
            if k.startswith('HATCH|A-DETL-GENF|SOLID') or k.startswith('LINHA|EC-LINHAS'):
                AVAC_COR[(tipo_avac(rel), k.split('|')[-1])].add(v['descricao'])
AVAC_COR = {c: next(iter(v)) for c, v in AVAC_COR.items() if len(v) == 1}
for rel, d in dados.items():
    if rel.startswith('AVAC/'):
        for e in d['entidades']:
            if '-CNTR' not in e['layer']:
                continue
            chave_cor = (tipo_avac(rel), e['cor'])
            if e['match'] == 'cor' and chave_cor not in AVAC_COR:
                e['descricao'], e['match'] = '', ''          # cor de outro sistema: anula
            if not e['descricao'] and chave_cor in AVAC_COR:
                e['descricao'], e['match'] = AVAC_COR[chave_cor], 'cor*'

wb = Workbook()
ws0 = wb.active
ws0.title = 'LEIA-ME'

# ================= Rastreio =================
wr = wb.create_sheet('Rastreio')
cols_r = ['ID', 'Pasta', 'Desenho', 'Piso', 'Família', 'Handle CAD', 'Tipo', 'Layer', 'Cor', 'Chave', 'Descrição (legenda)',
          'Match', 'Comp. (m)', 'Área (m²)', 'Un', 'X (m)', 'Y (m)']
cab(wr, 1, cols_r, [8, 22, 34, 14, 8, 10, 11, 16, 10, 44, 60, 7, 10, 10, 5, 10, 10])
linha = 2
agreg = collections.OrderedDict()   # (fam, rel, chave) -> info
sem_legenda = []
blocos = collections.Counter()
for rel, d in dados.items():
    fam = familia(rel)
    pasta, des = rel.rsplit('/', 1)
    piso = piso_de(rel)
    leg_por_cor = collections.defaultdict(set)
    for k, v in d['legenda'].items():
        leg_por_cor[k.split('|')[-1]].add(v['descricao'])
    # todas as entradas de legenda entram nas folhas (mesmo com quantidade 0)
    for k, v in d['legenda'].items():
        if fam in ('PAV', 'TET', 'PAR', 'AVAC') and not (fam == 'AVAC' and k.startswith('BLOCO')):
            un = 'un' if k.startswith('BLOCO') else ('m²' if k.startswith('HATCH') else 'm')
            if fam == 'AVAC' and k.startswith('HATCH'):
                continue  # amostras de cor; medição AVAC é feita pelas linhas de eixo (CNTR)
            agreg.setdefault((fam, rel, k), {'desc': v['descricao'], 'un': un, 'match': 'chave'})
    nm_ag = collections.defaultdict(lambda: [0, 0.0, 0.0])
    for e in d['entidades']:
        if e['legenda']:
            continue
        k = e['chave']
        incluir = False
        un = ''
        if k.startswith('BLOCO'):
            incluir, un = True, 'un'
            blocos[(rel, e['padrao_bloco'])] += 1
            if fam == 'INF' and re.match(r'(INFRA-|Incêndios|tomada eletrica)', e['padrao_bloco']):
                agreg.setdefault(('INF', rel, k), {'desc': e['padrao_bloco'].split(' - ')[0].strip(), 'un': 'un', 'match': 'bloco'})
        elif e['descricao'] and fam in ('PAV', 'TET', 'PAR', 'INF'):
            if fam == 'INF':
                e = dict(e, match=e['match'] + '?')
            incluir = True
            un = 'm²' if e['tipo'] == 'HATCH' else 'm'
        elif fam == 'AVAC' and '-CNTR' in e['layer'] and e['tipo'] != 'HATCH':
            incluir, un = True, 'm'
        if not incluir:
            if fam in ('PAV', 'TET', 'PAR') and not e['descricao']:
                a = nm_ag[k]; a[0] += 1; a[1] += e['comp_m']; a[2] += e['area_m2']
            continue
        wr.append([linha - 1, pasta, des, piso, fam, e['handle'], e['tipo'], e['layer'], e['cor'], k, e['descricao'],
                   e['match'], round(e['comp_m'], 3), round(e['area_m2'], 3), 1 if un == 'un' else None,
                   round(e['x'], 2) if e['x'] is not None else None, round(e['y'], 2) if e['y'] is not None else None])
        linha += 1
        if fam in ('PAV', 'TET', 'PAR', 'AVAC'):
            if un == 'un' and fam == 'AVAC':
                continue
            if not (fam == 'PAV' and un == 'un') and not (fam in ('TET', 'PAR') and un == 'un'):
                key = (fam, rel, k)
                if key not in agreg:
                    desc = e['descricao'] or ('Eixo ' + e['layer'] + ' ' + e['cor'])
                    agreg[key] = {'desc': desc, 'un': un, 'match': e['match'] or 'layer'}
    # sem legenda (só acabamentos): quantidades relevantes
    def cinz(c):
        if not c.startswith('#'):
            return True
        r_, g_, b_ = int(c[1:3], 16), int(c[3:5], 16), int(c[5:7], 16)
        return max(r_, g_, b_) - min(r_, g_, b_) < 24
    for k, (n, comp, area) in nm_ag.items():
        cor_k = k.split('|')[-1]
        cands = leg_por_cor.get(cor_k, set())
        if k.startswith('BLOCO') or (area < 1 and comp < 5):
            continue
        if cinz(cor_k) and not cands:
            continue
        if k.split('|')[1] in ('arquitectura', 'estrutura') and k.startswith('LINHA'):
            continue
        sem_legenda.append([pasta, des, piso, fam, k, n, round(comp, 2), round(area, 2), ' | '.join(sorted(cands))[:400]])
ult_r = linha - 1
for row in wr.iter_rows(min_row=2, max_row=ult_r):
    for c in row:
        c.font = fnt
    for c in row[12:14]:
        c.number_format = '#,##0.000'
tabela(wr, f'A1:Q{ult_r}', 'Rastreio')
R = "Rastreio!"

# ================= folhas por família =================
def folha_familia(nome, fam, unid_padrao, titulo, nome_tab, extra_pd=False):
    ws = wb.create_sheet(nome)
    ws['A1'] = titulo
    ws['A1'].font = ft
    ws['A2'] = ('Quantidade = SUMIFS sobre o separador Rastreio (Desenho + Chave). Interior/Exterior é editável '
                '(usado na verificação pavimento × teto).') if fam == 'PAV' else 'Quantidade = SUMIFS sobre o separador Rastreio (Desenho + Chave).'
    ws['A2'].font = Font(name=F, size=9, italic=True)
    hdr = 4
    cols = ['Pasta', 'Desenho', 'Piso', 'Descrição (legenda do desenho)', 'Chave CAD (tipo|layer|padrão|cor)', 'Un', 'Quantidade', 'Nº elementos', 'Match']
    widths = [22, 30, 14, 70, 44, 5, 12, 10, 8]
    if fam == 'PAV':
        cols.append('Int/Ext'); widths.append(8)
    if extra_pd:
        cols += ['Pé-direito (m)', 'Área (m²) = m × PD']; widths += [11, 14]
    cab(ws, hdr, cols, widths)
    r = hdr + 1
    for (fm, rel, k), inf in agreg.items():
        if fm != fam:
            continue
        pasta, des = rel.rsplit('/', 1)
        un = inf['un']
        ws.cell(r, 1, pasta); ws.cell(r, 2, des); ws.cell(r, 3, piso_de(rel)); ws.cell(r, 4, inf['desc']); ws.cell(r, 5, k); ws.cell(r, 6, un)
        if un == 'm²':
            ws.cell(r, 7, f'=SUMIFS({R}$N:$N,{R}$C:$C,$B{r},{R}$B:$B,$A{r},{R}$J:$J,$E{r})')
        elif un == 'un':
            ws.cell(r, 7, f'=COUNTIFS({R}$C:$C,$B{r},{R}$B:$B,$A{r},{R}$J:$J,$E{r})')
        else:
            ws.cell(r, 7, f'=SUMIFS({R}$M:$M,{R}$C:$C,$B{r},{R}$B:$B,$A{r},{R}$J:$J,$E{r})')
        ws.cell(r, 8, f'=COUNTIFS({R}$C:$C,$B{r},{R}$B:$B,$A{r},{R}$J:$J,$E{r})')
        ws.cell(r, 9, inf['match'])
        c = 10
        if fam == 'PAV':
            ws.cell(r, c, 'Ext' if EXTERIOR.search(inf['desc']) else 'Int').font = Font(name=F, size=10, color='0000FF'); c += 1
        if extra_pd:
            pd = ws.cell(r, c, f"=IFERROR(INDEX($M$2:$S$2,MATCH($C{r},$M$1:$S$1,0)),\"\")")
            pd.font = fln
            ws.cell(r, c + 1, f'=IF(AND(ISNUMBER(J{r}),$F{r}="m"),G{r}*J{r},"")').number_format = '#,##0.00'
        for cc in range(1, len(cols) + 1):
            x = ws.cell(r, cc)
            if x.font != fln and x.font.color is None or x.font == Font():
                x.font = fnt
            x.border = brd
        ws.cell(r, 4).alignment = wrap
        ws.cell(r, 7).number_format = '#,##0.00'
        r += 1
    last = r - 1
    if last > hdr:
        ws.conditional_formatting.add(f'G{hdr+1}:G{last}', CellIsRule(operator='equal', formula=['0'], fill=fill_warn))
    ws.auto_filter.ref = f'A{hdr}:{get_column_letter(len(cols))}{max(last, hdr+1)}'
    if extra_pd:
        ws['L1'] = 'Pé-direito por piso (m) ►'; ws['L1'].font = fb
        ws['L2'] = 'PREENCHER (amarelo)'; ws['L2'].font = Font(name=F, size=9, italic=True)
        for i, p in enumerate(['Piso 0', 'Piso 1', 'Piso 2', 'Piso 3', 'Cobertura', 'Ruína (P1/P2)', 'Exterior']):
            a = ws.cell(1, 13 + i, p); a.font = fb; a.fill = fill_sub; a.border = brd
            b = ws.cell(2, 13 + i, None); b.fill = fill_in; b.font = fin; b.border = brd; b.number_format = '0.00'
            ws.column_dimensions[get_column_letter(13 + i)].width = 11
        ws['M2'] = 2.70
        ws['M3'] = 'Exemplo: 2,70 (valor ilustrativo, substituir pelo pé-direito de projeto)'
        ws['M3'].font = Font(name=F, size=8, italic=True, color='808080')
    return ws, hdr, last


wp, hp, lp = folha_familia('Pavimentos', 'PAV', 'm²', 'Pavimentos — áreas por acabamento (hatches identificados pela legenda)', 'Pav')
wt, ht, lt = folha_familia('Tetos', 'TET', 'm²', 'Tetos — áreas por tipo de teto (hatches identificados pela legenda)', 'Tet')
wpa, hpa, lpa = folha_familia('Paredes', 'PAR', 'm', 'Paredes — desenvolvimento (m) por acabamento, medido em planta', 'Par', extra_pd=True)
wi, hi, li = folha_familia('Infraestruturas', 'INF', 'm', 'Infraestruturas (ARQ 02) — elementos lineares/áreas identificados pela legenda', 'Inf')
wa, ha, la = folha_familia('AVAC', 'AVAC', 'm', 'AVAC — comprimentos pelas linhas de eixo (layers *-CNTR), por sistema', 'Avac')

# ================= Blocos (contagens) =================
wbk = wb.create_sheet('Blocos')
wbk['A1'] = 'Contagem de blocos (símbolos / famílias Revit) por desenho'
wbk['A1'].font = ft
wbk['A2'] = 'Quantidade = COUNTIFS sobre Rastreio (Desenho + Chave BLOCO|família). Filtrar por Família/Desenho.'
wbk['A2'].font = Font(name=F, size=9, italic=True)
cab(wbk, 4, ['Pasta', 'Desenho', 'Piso', 'Família', 'Bloco (família - tipo)', 'Chave', 'Quantidade (un)'], [22, 34, 14, 8, 70, 30, 12])
r = 5
for (rel, nome), n in sorted(blocos.items()):
    pasta, des = rel.rsplit('/', 1)
    k = 'BLOCO|' + nome
    vals = [pasta, des, piso_de(rel), familia(rel), nome, k, f'=COUNTIFS({R}$C:$C,$B{r},{R}$B:$B,$A{r},{R}$J:$J,$F{r})']
    for i, v in enumerate(vals, 1):
        c = wbk.cell(r, i, v); c.font = fnt; c.border = brd
    r += 1
wbk.auto_filter.ref = f'A4:G{r-1}'

# ================= Resumo =================
wsr = wb.create_sheet('Resumo', 1)
wsr['A1'] = 'N2408 — Hostel Bonfim · Resumo de medições extraídas dos DWG'
wsr['A1'].font = ft
wsr['A2'] = 'Totais por descrição × piso. Fórmulas SUMIFS sobre os separadores de família (que por sua vez somam o Rastreio).'
wsr['A2'].font = Font(name=F, size=9, italic=True)
r = 4
blocos_resumo = [('Pavimentos', 'PAV', 'm²', hp, lp), ('Tetos', 'TET', 'm²', ht, lt), ('Paredes', 'PAR', 'm', hpa, lpa),
                 ('Infraestruturas', 'INF', None, hi, li), ('AVAC', 'AVAC', 'm', ha, la)]
for nome, fam, un, h0, h1 in blocos_resumo:
    c = wsr.cell(r, 1, f'{nome}'); c.font = Font(name=F, size=12, bold=True, color='1F3864')
    r += 1
    hdrs = ['Descrição', 'Un'] + PISOS + ['Total']
    for i, t in enumerate(hdrs, 1):
        x = wsr.cell(r, i, t); x.font = fh; x.fill = fill_h; x.border = brd; x.alignment = Alignment(wrap_text=True)
    r += 1
    descs = collections.OrderedDict()
    for (fm, rel, k), inf in agreg.items():
        if fm == fam:
            descs.setdefault((inf['desc'], inf['un']), None)
    ini = r
    for (desc, u) in sorted(descs, key=lambda x: (x[1], x[0])):
        wsr.cell(r, 1, desc).alignment = wrap
        wsr.cell(r, 2, u)
        for j, p in enumerate(PISOS):
            col = get_column_letter(3 + j)
            rng = f"'{nome}'!"
            f = (f"=SUMIFS({rng}$G${h0+1}:$G${h1},{rng}$D${h0+1}:$D${h1},$A{r},{rng}$F${h0+1}:$F${h1},$B{r},"
                 f"{rng}$C${h0+1}:$C${h1},\"{p}*\")")
            wsr.cell(r, 3 + j, f).number_format = '#,##0.00;-#,##0.00;"-"'
        wsr.cell(r, 3 + len(PISOS), f'=SUM(C{r}:{get_column_letter(2+len(PISOS))}{r})').number_format = '#,##0.00;-#,##0.00;"-"'
        for cc in range(1, 4 + len(PISOS)):
            x = wsr.cell(r, cc); x.border = brd
            if x.font == Font() or x.font.name != F:
                x.font = fln if cc >= 3 and cc < 3 + len(PISOS) else fnt
        wsr.cell(r, 3 + len(PISOS)).font = fb
        r += 1
    r += 1
wsr.column_dimensions['A'].width = 80
wsr.column_dimensions['B'].width = 5
for j in range(len(PISOS) + 1):
    wsr.column_dimensions[get_column_letter(3 + j)].width = 12
wsr.freeze_panes = 'C4'

# ================= Sem legenda =================
wsl = wb.create_sheet('Sem legenda')
wsl['A1'] = 'Elementos de acabamento desenhados SEM correspondência na legenda — validar'
wsl['A1'].font = ft
wsl['A2'] = ('Cor/padrão que não coincide com nenhuma amostra da legenda do próprio desenho. "Candidatos" = itens da legenda com a mesma cor. '
             'Possível erro de desenho/legenda (incompatibilidade) ou zona sem intervenção.')
wsl['A2'].font = Font(name=F, size=9, italic=True)
cab(wsl, 4, ['Pasta', 'Desenho', 'Piso', 'Família', 'Chave CAD', 'Nº elementos', 'Comp. (m)', 'Área (m²)', 'Candidatos na legenda (mesma cor)', 'Decisão (preencher)'],
    [22, 14, 12, 7, 40, 9, 10, 10, 70, 30])
r = 5
for row in sorted(sem_legenda, key=lambda x: (x[1], -x[7], -x[6])):
    for i, v in enumerate(row + [None], 1):
        c = wsl.cell(r, i, v); c.font = fin if i in (6, 7, 8) else fnt; c.border = brd
    wsl.cell(r, 9).alignment = wrap
    wsl.cell(r, 10).fill = fill_in
    r += 1

# ================= Verificações (compatibilização) =================
wv = wb.create_sheet('Verificações')
wv['A1'] = 'Verificações de coerência e incompatibilidades detetadas'
wv['A1'].font = ft
cab(wv, 3, ['Piso', 'Pavimento interior (m²)', 'Teto (m²)', 'Diferença (m²)', 'Dif. %', 'Observação'], [16, 18, 14, 14, 9, 70])
r = 4
for p in ['Piso 0', 'Piso 1', 'Piso 2', 'Piso 3']:
    wv.cell(r, 1, p)
    wv.cell(r, 2, f"=SUMIFS(Pavimentos!$G${hp+1}:$G${lp},Pavimentos!$C${hp+1}:$C${lp},$A{r},Pavimentos!$F${hp+1}:$F${lp},\"m²\",Pavimentos!$J${hp+1}:$J${lp},\"Int\")")
    wv.cell(r, 3, f"=SUMIFS(Tetos!$G${ht+1}:$G${lt},Tetos!$C${ht+1}:$C${lt},$A{r},Tetos!$F${ht+1}:$F${lt},\"m²\")")
    wv.cell(r, 4, f'=B{r}-C{r}')
    wv.cell(r, 5, f'=IF(C{r}=0,"",D{r}/C{r})')
    wv.cell(r, 6, f'=IF(C{r}=0,"",IF(ABS(E{r})>0.1,"Diferença >10% — verificar zonas sem legenda / sem teto / escadas","OK (≤10%)"))')
    for cc in range(1, 7):
        x = wv.cell(r, cc); x.font = fln if cc in (2, 3) else fnt; x.border = brd
    for cc in (2, 3, 4):
        wv.cell(r, cc).number_format = '#,##0.00'
    wv.cell(r, 5).number_format = '0.0%'
    r += 1
wv.cell(r + 1, 1, 'Nota: "Pavimento interior" exclui itens marcados Ext na coluna Int/Ext do separador Pavimentos (editável).').font = Font(name=F, size=9, italic=True)

r += 3
wv.cell(r, 1, 'Incompatibilidades / observações').font = Font(name=F, size=12, bold=True, color='1F3864')
r += 1
for i, t in enumerate(['#', 'Tipo', 'Desenhos', 'Descrição', 'Impacto / ação sugerida', 'Estado'], 1):
    x = wv.cell(r, i, t); x.font = fh; x.fill = fill_h; x.border = brd
r += 1
OBS = [
    ('Nomenclatura', 'ARQ (1-xx, 3-xx) vs AVAC (101/201)', 'ARQ designa os pisos 0/1/2/3; AVAC designa -1/1/2/3 (não existe "Piso 0" no AVAC). Assumido Piso -1 AVAC = Piso 0 ARQ.', 'Confirmar com projetistas e uniformizar a designação de pisos em todas as especialidades.'),
    ('Numeração', 'ARQ/01-Desenhos Gerais 2-06…2-11 vs ARQ/02-Infraestruturas 2-06…2-11', 'Mesma numeração de desenho para conteúdos diferentes (alçados/cortes vs plantas de infraestruturas).', 'Pedir lista de peças desenhadas corrigida; risco de troca de desenhos em obra.'),
    ('Versões', 'ARQ/07 N2408-ARQ-Q-B-3-06 e "…3-06 v2"', 'Duas versões do mesmo desenho de quarto de banho na mesma pasta.', 'Confirmar qual está em vigor e retirar a outra.'),
    ('Projeto incompleto', 'ARQ/03 Tetos 3-11 (legenda)', 'Nota do projetista: "Tetos do piso 3 têm de ser confirmados em obra e se possível aumentar a sua altura".', 'Quantidades de tetos do Piso 3 sujeitas a revisão.'),
    ('Legenda', 'Pavimentos 3-01…3-04', 'Hatches com cor/padrão sem amostra correspondente na legenda (ver separador "Sem legenda").', 'Esclarecer acabamento destas áreas com o projetista.'),
    ('Legenda', 'Pavimentos (legenda)', '3 entradas "Nordik Arce 20x120" com o mesmo texto base e cor; diferem só no suporte (Acoustikcork / betonilha c/ manta acústica / sem indicação).', 'Confirmar o suporte por compartimento — impacto em custo.'),
    ('Ficheiros não processados', 'Ver separador Ficheiros (estado "NÃO PROCESSADO")', 'DWG > 5 MB não puderam ser descarregados pela ligação ao Drive: Vãos Interiores, Mapa de Armários, AVAC 102/103/104 (aeráulicas pisos 1-3), AVAC 301/302 (isométricas), SCIE.', 'Medições AVAC de condutas dos pisos 1, 2 e 3, vãos interiores, armários e SCIE em falta.'),
    ('Legenda', 'Pavimentos 3-01 (Piso 0)', 'Zona poente: ~165 m² de hatches cor #FF7777 (handles 3C57, 3E96, 3D52, …) sobrepostos aos acabamentos da legenda, sem item próprio na legenda.', 'Esclarecer o significado (existente? demolição? reforço?) — ver verificacao/3-01_verificacao.png.'),
    ('Legenda', 'Pavimentos 3-02 (Piso 1)', 'Pavimento exterior envolvente (~305 m², cinzento #C0C0C0, handles 9C77, 9CDA, 98D1) sem amostra igual; 2 itens da legenda com a mesma cor (cubo de granito existente / borracha). Rampa/zona a vermelho (B5D8, 21 m²) sem legenda.', 'Confirmar o acabamento exterior com o projetista — ver verificacao/3-02_verificacao.png.'),
    ('Coerência entre plantas', 'Piso 1: 1-02, 2-07, 3-02, 3-06, 3-11, 5-05', 'Beliches: 21 nas plantas geral/infra vs 24 nas de pavimentos/tetos/vãos; sanitas 14 vs 15; bases de duche 10 vs 11; armários 20–24; janelas 42–72 (vistas com recortes diferentes).', 'Indica plantas não atualizadas entre si — confirmar com o projetista qual a versão correta (impacta mobiliário, louças e vãos).'),
    ('Coerência entre plantas', 'Piso 2: 1-03, 2-08, 3-03, 3-07, 3-12, 5-06', 'Portas interiores: 65–66 nas plantas ARQ vs 56 na planta de vãos 5-06; beliches 44 vs 47; armários 11 vs 14; sanitas/duches/lavatórios +1 nas plantas de pavimento e tetos.', 'Mapa de vãos pode estar incompleto — validar contagem de portas antes de orçamentar.'),
    ('Coerência entre plantas', 'Piso 3: 1-04, 2-09, 3-04, 3-08, 3-13', 'Armários quarto: 9–13 conforme a planta; móveis lavatório 7 vs 8.', 'Confirmar quantidades de mobiliário fixo no piso 3.'),
    ('AVAC', 'AVAC 201/202/204 (tubagem)', 'Linhas de eixo P-PIPE-CNTR a vermelho (#FF0000/#CC0000), ~23 m no total, sem item correspondente na legenda.', 'Pedir ao projetista AVAC a identificação deste sistema (possível rede de condensados).'),
    ('Método', 'Paredes 3-05…3-09', 'Acabamentos de parede desenhados como linhas em planta: mede-se o desenvolvimento (m). Área = m × pé-direito (preencher no separador Paredes).', 'Validar pé-direito por piso (ver alturas de teto na legenda de tetos).'),
]
for i, (tp, ds, de, ac) in enumerate(OBS, 1):
    for j, v in enumerate([i, tp, ds, de, ac, None], 1):
        c = wv.cell(r, j, v); c.font = fnt; c.border = brd; c.alignment = wrap
    wv.cell(r, 6).fill = fill_in
    r += 1
wv.column_dimensions['B'].width = 18; wv.column_dimensions['C'].width = 30; wv.column_dimensions['D'].width = 70; wv.column_dimensions['E'].width = 55

# --- comparação de blocos entre desenhos do mesmo piso ---
r += 2
wv.cell(r, 1, 'Coerência entre desenhos do mesmo piso — contagem de blocos (COUNTIFS sobre Rastreio)').font = Font(name=F, size=12, bold=True, color='1F3864')
r += 1
wv.cell(r, 1, 'Se o mesmo elemento aparece em quantidades diferentes nas plantas do mesmo piso, alguma planta não foi atualizada (ou o elemento está omitido nessa vista).').font = Font(name=F, size=9, italic=True)
r += 1
DESENHOS = {  # piso -> [(rotulo, pasta, desenho)]
    'Piso 0': [('Geral', 'ARQ/01-Desenhos Gerais', '1-01'), ('Infra', 'ARQ/02-Infraestruturas', '2-06'), ('Pav', 'ARQ/03-Acabamentos/Pavimentos', '3-01'),
               ('Par', 'ARQ/03-Acabamentos/Paredes', '3-05'), ('Tet', 'ARQ/03-Acabamentos/Tetos', '3-10'), ('Vãos', None, None)],
    'Piso 1': [('Geral', 'ARQ/01-Desenhos Gerais', '1-02'), ('Infra', 'ARQ/02-Infraestruturas', '2-07'), ('Pav', 'ARQ/03-Acabamentos/Pavimentos', '3-02'),
               ('Par', 'ARQ/03-Acabamentos/Paredes', '3-06'), ('Tet', 'ARQ/03-Acabamentos/Tetos', '3-11'), ('Vãos', 'ARQ/05-Vãos Exteriores', '5-05')],
    'Piso 2': [('Geral', 'ARQ/01-Desenhos Gerais', '1-03'), ('Infra', 'ARQ/02-Infraestruturas', '2-08'), ('Pav', 'ARQ/03-Acabamentos/Pavimentos', '3-03'),
               ('Par', 'ARQ/03-Acabamentos/Paredes', '3-07'), ('Tet', 'ARQ/03-Acabamentos/Tetos', '3-12'), ('Vãos', 'ARQ/05-Vãos Exteriores', '5-06')],
    'Piso 3': [('Geral', 'ARQ/01-Desenhos Gerais', '1-04'), ('Infra', 'ARQ/02-Infraestruturas', '2-09'), ('Pav', 'ARQ/03-Acabamentos/Pavimentos', '3-04'),
               ('Par', 'ARQ/03-Acabamentos/Paredes', '3-08'), ('Tet', 'ARQ/03-Acabamentos/Tetos', '3-13'), ('Vãos', None, None)],
}
CATS = [('Portas interiores (1 folha + correr)', ['BLOCO|porta interior*', 'BLOCO|Porta correr*']), ('Janelas', ['BLOCO|*janela*']),
        ('Sanitas', ['BLOCO|WC_sanita*']), ('Bases de duche', ['BLOCO|WC_base_duche*']), ('Móveis lavatório', ['BLOCO|*lavatório*']),
        ('Beliches', ['BLOCO|QUA_Beliche*']), ('Armários quarto', ['BLOCO|QUA_armário*'])]
hd = ['Piso', 'Elemento', '1-0x Geral', '2-0x Infra', '3-0x Pav', '3-0x Par', '3-1x Tet', '5-0x Vãos', 'Máx − Mín', 'Resultado']
for i, t in enumerate(hd, 1):
    x = wv.cell(r, i, t); x.font = fh; x.fill = fill_h; x.border = brd; x.alignment = Alignment(wrap_text=True)
r += 1
ini_cmp = r
for piso, lst in DESENHOS.items():
    for cat, pats in CATS:
        wv.cell(r, 1, piso); wv.cell(r, 2, cat)
        for j, (rot, pasta, des) in enumerate(lst):
            if not pasta or f'{pasta}/{des}' not in dados:
                wv.cell(r, 3 + j, 'n/d')
                continue
            f = '+'.join(f'COUNTIFS({R}$B:$B,"{pasta}",{R}$C:$C,"{des}",{R}$J:$J,"{pt}")' for pt in pats)
            wv.cell(r, 3 + j, '=' + f)
        wv.cell(r, 9, f'=MAX(C{r}:H{r})-MIN(C{r}:H{r})')
        wv.cell(r, 10, f'=IF(MAX(C{r}:H{r})=0,"—",IF(I{r}=0,"OK","DIFERENÇA — verificar"))')
        for cc in range(1, 11):
            x = wv.cell(r, cc); x.font = fnt; x.border = brd
        r += 1
wv.conditional_formatting.add(f'J{ini_cmp}:J{r-1}', CellIsRule(operator='equal', formula=['"DIFERENÇA — verificar"'], fill=fill_warn))
wv.cell(r, 1, 'Nota: uma vista pode omitir propositadamente equipamento (ex.: plantas de tetos); interpretar as diferenças com essa ressalva.').font = Font(name=F, size=9, italic=True)

# ================= Ficheiros =================
wf = wb.create_sheet('Ficheiros')
wf['A1'] = 'Inventário de ficheiros (Google Drive › Hostel Bonfim › 01_Projectos)'
wf['A1'].font = ft
cab(wf, 3, ['Pasta', 'Ficheiro', 'Piso', 'Família', 'Tamanho (KB)', 'Unidades DWG', 'Entidades medidas', 'Itens de legenda', 'Estado', 'Link Drive'],
    [24, 55, 14, 9, 11, 11, 12, 11, 18, 50])
UN = {1: 'pol', 4: 'mm', 5: 'cm', 6: 'm'}
r = 4
for rel, m in sorted(man.items()):
    d = dados.get(rel)
    estado = 'Processado' if d else 'NÃO PROCESSADO (>5 MB)'
    vals = [m['pasta'], m['ficheiro'], piso_de(rel), familia(rel), round(int(m['tamanho']) / 1024), UN.get(d['unidades'], d['unidades']) if d else '-',
            len(d['entidades']) if d else 0, len(d['legenda']) if d else 0, estado, f"https://drive.google.com/file/d/{m['id']}/view"]
    for i, v in enumerate(vals, 1):
        c = wf.cell(r, i, v); c.font = fnt; c.border = brd
    if not d:
        for i in range(1, 11):
            wf.cell(r, i).fill = fill_warn
    r += 1
wf.auto_filter.ref = f'A3:J{r-1}'

# ================= LEIA-ME =================
L = [
    ('N2408 — Hostel Bonfim · Medições extraídas dos DWG', ft),
    (f'Gerado em {datetime.date.today():%d/%m/%Y} a partir dos DWG no Google Drive (pasta 01_Projectos: ARQ, AVAC, SCIE).', fnt),
    ('', fnt),
    ('Como foi feito (passos rastreáveis)', fb),
    ('1. Descarga dos DWG do Google Drive (ID de cada ficheiro no separador Ficheiros).', fnt),
    ('2. Conversão DWG → DXF com CADCraft (open source, github.com/storytold/cadcraft); 0 entidades perdidas na conversão.', fnt),
    ('3. Leitura de cada entidade do espaço modelo: comprimento (linhas, polilinhas, arcos), área (hatches, com ilhas descontadas), contagem (blocos). Unidades normalizadas para metros.', fnt),
    ('4. Identificação automática da legenda de cada desenho (amostra + texto à direita) e classificação de cada elemento pela "chave" (tipo|layer|padrão|cor) igual à da amostra.', fnt),
    ('5. Separador Rastreio: 1 linha por elemento CAD com o Handle — permite localizar o elemento no AutoCAD (comando _SELECT / Quick Select por handle).', fnt),
    ('6. Separadores de família somam o Rastreio (SUMIFS/COUNTIFS); o Resumo soma os separadores de família por piso.', fnt),
    ('', fnt),
    ('Separadores', fb),
    ('Resumo · Pavimentos · Tetos · Paredes · Infraestruturas · AVAC · Blocos · Sem legenda · Verificações · Ficheiros · Rastreio', fnt),
    ('', fnt),
    ('Convenções de cor', fb),
    ('Texto azul = dado extraído do CAD / valor de entrada · Texto verde = ligação a outro separador · Texto preto = fórmula', fnt),
    ('Fundo amarelo = célula a preencher pelo utilizador (pé-direito, decisões, estado) · Fundo salmão = quantidade 0 ou ficheiro não processado', fnt),
    ('Match: "chave" = igual à amostra da legenda · "cor" = só a cor coincide (cor única na legenda) → validar · "cor*" = cor da legenda de outro desenho AVAC · "layer" = eixo AVAC sem legenda · "bloco" = contagem de blocos · "?" = legenda de infraestruturas pouco fiável', fnt),
    ('', fnt),
    ('Limitações', fb),
    ('• Mede o que está desenhado: hatches sobrepostos, zonas não hachuradas ou elementos fora de escala não são corrigidos automaticamente.', fnt),
    ('• Paredes: mede-se o desenvolvimento em planta (m). A área exige o pé-direito (amarelo no separador Paredes; 2,70 m é apenas exemplo).', fnt),
    ('• AVAC: comprimento pelas linhas de eixo (*-CNTR). Diâmetros/secções não desagregados nesta versão. Pisos 1–3 de condutas não processados (ficheiros > 5 MB).', fnt),
    ('• Infraestruturas: contagens por família de bloco no separador Blocos (ex.: INFRA-tomada); linhas da legenda no separador Infraestruturas.', fnt),
    ('• Atribuição de pisos assumida a partir do nome das vistas Revit (ex.: "03_ACAB PAV - piso 1"). 1-01 e 2-06 (sem nome de vista) assumidos como Piso 0.', fnt),
]
for i, (t, f_) in enumerate(L, 1):
    c = ws0.cell(i, 1, t); c.font = f_
ws0.column_dimensions['A'].width = 150

wb.move_sheet('Rastreio', offset=len(wb.sheetnames))
wb.save(SAIDA)
print('ok', SAIDA, 'rastreio', ult_r - 1, 'agreg', len(agreg), 'sem_legenda', len(sem_legenda), 'blocos', len(blocos))
