import sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

sp, out = sys.argv[1], sys.argv[2]
F = "Arial"
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="1F3864"); HFont = Font(name=F, bold=True, color="FFFFFF", size=10)
INPUT = PatternFill("solid", fgColor="FFF2CC"); BODY = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color="0000FF"); GREEN = Font(name=F, size=10, color="008000")
EUR = '#,##0.00 "€";-#,##0.00 "€";-'
ESP = ["GER", "ARQ", "EST", "HID", "PLU", "ESG", "ELE", "ITED", "AVAC", "GAS", "SCIE", "TERM", "ACU", "CARP"]
ESP_NOME = {"GER": "Gerais / Estaleiro", "ARQ": "Arquitetura", "EST": "Estruturas", "HID": "Águas (abastecimento)", "PLU": "Águas pluviais",
            "ESG": "Esgotos", "ELE": "Eletricidade", "ITED": "ITED", "AVAC": "AVAC / Ventilação", "GAS": "Gás", "SCIE": "SCIE",
            "TERM": "Térmica", "ACU": "Acústica", "CARP": "Carpintarias / Mobiliário"}
PREFIX = {"GER": "GER", "ARQ": "ARQ", "EST": "EST", "HID": "HID", "PLU": "PLU", "ESG": "ESG", "ELE": "ELE", "ITED": "ITE",
          "AVAC": "AVA", "GAS": "GAS", "SCIE": "SCI", "TERM": "TER", "ACU": "ACU", "CARP": "CAR"}

rows = []; maxn = defaultdict(int)
for l in open(f"{sp}/base_mercado.txt", encoding="utf-8"):
    c, d, e, u, p = l.strip().split("|")
    pre, n = c.split("."); maxn[pre] = max(maxn[pre], int(n))
    obs = "Código mantido da base 2026; especialidade reclassificada" if (c.startswith("HID.0") and e == "ESG") or (c.startswith("TER.0") and e == "ACU") else ""
    rows.append(dict(c=c, d=d, e=e, u=u, glob=float(p), fonte="Referência de mercado Porto/Aveiro (BASE_PRECOS_2026)", obra="", data="Jun-2026", fiab="C", obs=obs))
for l in open(f"{sp}/base_ribeira.txt", encoding="utf-8"):
    e, d, u, p = l.strip().split("|")
    pre = PREFIX[e]; maxn[pre] += 1
    rows.append(dict(c=f"{pre}.{maxn[pre]:02d}", d=d, e=e, u=u, glob=float(p), fonte="Orçamento próprio (custo, antes de margem)", obra="Ribeira – remodelação (L99)",
                     data="2026", fiab="B", obs="Inclui agravamento por acesso condicionado (transporte manual 100 m + escadaria, 1.º andar s/ elevador)"))
order = {e: i for i, e in enumerate(ESP)}
rows.sort(key=lambda r: (order[r["e"]], r["c"]))

wb = Workbook()
# ---------------- Instruções
ins = wb.active; ins.title = "Instruções"; ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 2; ins.column_dimensions["B"].width = 26; ins.column_dimensions["C"].width = 100
ins["B1"] = "Base de Preços PROBUILT – preços reais"; ins["B1"].font = Font(name=F, bold=True, size=16, color="1F3864")
ins["B2"] = "v1.0 · Margem calculada SOBRE O CUSTO · Estrutura compatível com a folha _Produção_Folha_de_orçamento.xlsx"
ins["B2"].font = Font(name=F, italic=True, size=10, color="595959")
r = 4
def sec(t):
    global r; ins.cell(r, 2, t).font = Font(name=F, bold=True, size=11, color="1F3864"); r += 1
def line(a, b, fill=None, font=None):
    global r
    x = ins.cell(r, 2, a); x.font = Font(name=F, bold=True, size=10); x.alignment = Alignment(wrap_text=True, vertical="top")
    if fill: x.fill = fill
    if font: x.font = font
    y = ins.cell(r, 3, b); y.font = BODY; y.alignment = Alignment(wrap_text=True, vertical="top"); r += 1
sec("COMO FUNCIONA O PREÇO")
line("1. Custo direto", "Soma de Mão de obra + Materiais + Equipamento + Subempreitada (colunas E–H). Se o artigo ainda não estiver decomposto, usa o «Preço global» (coluna I).")
line("2. Custo c/ encargos", "Cada parcela é agravada pela percentagem de encargos da folha Config (iguais à sua folha de orçamento: MO 1 %, Mat. 3 %, Equip. 2 %, Subemp. 2 %). O preço global usa o encargo médio.")
line("3. Preço de venda", "PV = Custo c/ encargos × (1 + Margem). A margem é SOBRE O CUSTO. Por omissão usa a margem padrão da Config; pode definir uma margem própria por artigo (coluna L).")
line("Exemplo", "Custo c/ encargos 100,00 € e margem 20 % → PV = 120,00 €. (A base 2026 anterior fazia 100 / 0,75 = 133,33 €, ou seja, 33 % sobre o custo.)")
r += 1
sec("FIABILIDADE DO PREÇO (coluna R)")
line("A", "Preço real: fatura, contrato ou adjudicação.", PatternFill("solid", fgColor="C6EFCE"))
line("B", "Preço de orçamento próprio ou proposta de subempreiteiro/fornecedor.", PatternFill("solid", fgColor="FFEB9C"))
line("C", "Referência de mercado – usar com cautela e substituir logo que haja preço B ou A.", PatternFill("solid", fgColor="FFC7CE"))
r += 1
sec("CÉLULAS A PREENCHER / CORES")
line("Fundo amarelo", "Dados de entrada (decomposição, preço global, margem própria, fonte, obra, data, fiabilidade).", INPUT)
line("Texto azul", "Valor introduzido manualmente (preço).", font=Font(name=F, bold=True, size=10, color="0000FF"))
line("Texto preto", "Fórmula – não editar.")
r += 1
sec("REGISTO DE PREÇOS (folha Registo)")
line("Para que serve", "Sempre que receber um preço novo (proposta, fatura, orçamento), registe-o com o código do artigo. A Base mostra automaticamente o n.º de registos e o último preço registado, para decidir quando atualizar o preço do artigo.")
line("Exemplo", "15/10/2026 · ARQ.21 · Moradia X · Proposta pintor Y · 7,90 € · B")
r += 1
sec("NOTAS")
for t in ["Origem: 115 artigos da BASE_PRECOS_2026 (referência de mercado, fiabilidade C) + 32 artigos do orçamento da Ribeira (L99), preços de custo antes de margem (fiabilidade B).",
          "Os preços da Ribeira incluem agravamento por acesso condicionado – ver observações; numa obra sem condicionantes podem ser mais baixos.",
          "HID.05–HID.08 passaram para ESG e TER.04–TER.05 para ACU (códigos mantidos para rastreabilidade).",
          "Especialidades alinhadas com as pastas do Drive: ARQ, EST, HID (inclui PLU), ESG, ELE, ITED, AVAC, GAS, SCIE, TERM, ACU, + GER e CARP para orçamento."]:
    c = ins.cell(r, 3, "• " + t); c.font = Font(name=F, size=9, color="595959"); c.alignment = Alignment(wrap_text=True); r += 1

# ---------------- Config
cf = wb.create_sheet("Config"); cf.sheet_view.showGridLines = False
cf.column_dimensions["A"].width = 2; cf.column_dimensions["B"].width = 36; cf.column_dimensions["C"].width = 12; cf.column_dimensions["D"].width = 60
cf["B1"] = "PARÂMETROS"; cf["B1"].font = Font(name=F, bold=True, size=14, color="1F3864")
params = [("Margem padrão (sobre o custo)", 0.20, "Igual à folha _Produção_Folha_de_orçamento.xlsx (célula I6)"),
          ("Encargos – Mão de obra", 0.01, "Folha de orçamento, célula J6"),
          ("Encargos – Materiais", 0.03, "Folha de orçamento, célula K6"),
          ("Encargos – Equipamento", 0.02, "Folha de orçamento, célula L6"),
          ("Encargos – Subempreitada", 0.02, "Folha de orçamento, célula M6"),
          ("Encargo médio – preço global", 0.02, "Pressuposto: média aplicada a artigos ainda não decompostos – ajustar"),
          ("IVA – taxa normal", 0.23, "Referência"),
          ("IVA – taxa reduzida (ARU / reabilitação)", 0.06, "Referência – confirmar enquadramento em cada obra")]
for i, (a, v, n) in enumerate(params):
    rr = 3 + i
    cf.cell(rr, 2, a).font = Font(name=F, bold=True, size=10)
    c = cf.cell(rr, 3, v); c.number_format = "0.0%"; c.fill = INPUT; c.font = BLUE; c.border = B; c.alignment = Alignment(horizontal="center")
    cf.cell(rr, 4, n).font = Font(name=F, italic=True, size=9, color="595959")
MARG, EMO, EMAT, EEQ, ESUB, EGLOB = ("Config!$C$3", "Config!$C$4", "Config!$C$5", "Config!$C$6", "Config!$C$7", "Config!$C$8")
cf["B13"] = "ESPECIALIDADES"; cf["B13"].font = Font(name=F, bold=True, size=11, color="1F3864")
for i, e in enumerate(ESP):
    cf.cell(14 + i, 2, e).font = Font(name=F, bold=True, size=10); cf.cell(14 + i, 3, ESP_NOME[e]).font = BODY
ESP_RANGE = f"=Config!$B$14:$B${13+len(ESP)}"

# ---------------- Registo
rg = wb.create_sheet("Registo")
rh = ["Data", "Cód.", "Obra", "Fonte (fornecedor / subempreiteiro / documento)", "Preço unitário (€)", "Fiabilidade", "Observações"]
rw = [12, 10, 28, 42, 15, 11, 45]
for i, (h, w) in enumerate(zip(rh, rw), 1):
    c = rg.cell(1, i, h); c.font = HFont; c.fill = HF; c.border = B; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rg.column_dimensions[get_column_letter(i)].width = w
for rr in range(2, 502):
    for j in range(1, 8):
        c = rg.cell(rr, j); c.fill = INPUT; c.border = B; c.font = BLUE if j == 5 else BODY
    rg.cell(rr, 1).number_format = "dd/mm/yyyy"; rg.cell(rr, 5).number_format = EUR
rg.freeze_panes = "A2"
dvf = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); rg.add_data_validation(dvf); dvf.add("F2:F501")

# ---------------- Base
bs = wb.create_sheet("Base", 1)
bs["A1"] = "BASE DE PREÇOS PROBUILT"; bs["A1"].font = Font(name=F, bold=True, size=14, color="1F3864")
bs["A2"] = '="Margem padrão sobre o custo: "&TEXT(Config!$C$3,"0%")&"   ·   Artigos: "&COUNTA(A5:A1000)&"   ·   Fiabilidade A/B/C: "&COUNTIF(R5:R1000,"A")&" / "&COUNTIF(R5:R1000,"B")&" / "&COUNTIF(R5:R1000,"C")'
bs["A2"].font = Font(name=F, italic=True, size=10, color="595959")
heads = ["Cód.", "Designação", "Esp.", "Un.", "Mão de obra (€)", "Materiais (€)", "Equipamento (€)", "Subempreitada (€)", "Preço global s/ decomp. (€)",
         "Custo direto (€)", "Custo c/ encargos (€)", "Margem própria (s/ custo)", "Margem aplicada", "PV unitário (€)",
         "Fonte", "Obra de origem", "Data", "Fiabilidade", "N.º registos", "Último preço registado (€)", "Observações"]
widths = [9, 52, 7, 6, 11, 11, 11, 12, 13, 12, 12, 11, 10, 12, 34, 24, 10, 10, 9, 13, 48]
HR = 4
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = bs.cell(HR, i, h); c.font = HFont; c.fill = HF; c.border = B; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bs.column_dimensions[get_column_letter(i)].width = w
bs.row_dimensions[HR].height = 42
EXTRA = 100
last = HR + len(rows) + EXTRA
for k in range(len(rows) + EXTRA):
    rr = HR + 1 + k
    r_ = rows[k] if k < len(rows) else None
    vals = [r_["c"], r_["d"], r_["e"], r_["u"]] if r_ else [None] * 4
    for j, v in enumerate(vals, 1):
        c = bs.cell(rr, j, v); c.font = BODY; c.border = B
        c.alignment = Alignment(horizontal="center", vertical="top") if j in (1, 3, 4) else Alignment(wrap_text=True, vertical="top")
        if not r_: c.fill = INPUT
    for j in range(5, 10):
        c = bs.cell(rr, j); c.fill = INPUT; c.border = B; c.font = BLUE; c.number_format = EUR
    if r_: bs.cell(rr, 9).value = r_["glob"]
    bs.cell(rr, 10, f'=IF(AND(A{rr}="",B{rr}=""),"",IF(SUM(E{rr}:H{rr})>0,SUM(E{rr}:H{rr}),I{rr}))')
    bs.cell(rr, 11, f'=IF(J{rr}="","",ROUND(IF(SUM(E{rr}:H{rr})>0,E{rr}*(1+{EMO})+F{rr}*(1+{EMAT})+G{rr}*(1+{EEQ})+H{rr}*(1+{ESUB}),I{rr}*(1+{EGLOB})),2))')
    c = bs.cell(rr, 12); c.fill = INPUT; c.border = B; c.font = BLUE; c.number_format = "0%"
    bs.cell(rr, 13, f'=IF(J{rr}="","",IF(L{rr}="",{MARG},L{rr}))')
    bs.cell(rr, 14, f'=IF(K{rr}="","",ROUND(K{rr}*(1+M{rr}),2))')
    for j, key in zip(range(15, 19), ["fonte", "obra", "data", "fiab"]):
        c = bs.cell(rr, j, r_[key] if r_ else None); c.fill = INPUT; c.border = B; c.font = BODY
        c.alignment = Alignment(horizontal="center", vertical="top") if j in (17, 18) else Alignment(wrap_text=True, vertical="top")
    bs.cell(rr, 19, f'=IF(A{rr}="","",COUNTIF(Registo!$B$2:$B$501,A{rr}))')
    bs.cell(rr, 20, f'=IF(OR(A{rr}="",S{rr}=0),"",LOOKUP(2,1/(Registo!$B$2:$B$501=A{rr}),Registo!$E$2:$E$501))')
    c = bs.cell(rr, 21, r_["obs"] if r_ else None); c.fill = INPUT; c.border = B; c.font = Font(name=F, size=9); c.alignment = Alignment(wrap_text=True, vertical="top")
    for j in (10, 11, 13, 14, 19, 20):
        c = bs.cell(rr, j); c.border = B; c.font = BODY; c.alignment = Alignment(horizontal="right", vertical="top")
    for j in (10, 11, 14, 20): bs.cell(rr, j).number_format = EUR
    bs.cell(rr, 13).number_format = "0%"
    bs.cell(rr, 14).font = Font(name=F, size=10, bold=True)
dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); bs.add_data_validation(dv); dv.add(f"R{HR+1}:R{last}")
dve = DataValidation(type="list", formula1=ESP_RANGE, allow_blank=True); bs.add_data_validation(dve); dve.add(f"C{HR+1}:C{last}")
for val, col in [("A", "C6EFCE"), ("B", "FFEB9C"), ("C", "FFC7CE")]:
    bs.conditional_formatting.add(f"R{HR+1}:R{last}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=col)))
bs.freeze_panes = f"C{HR+1}"; bs.auto_filter.ref = f"A{HR}:U{last}"
bs["L4"].comment = Comment("Deixe vazio para usar a margem padrão da folha Config. Margem SOBRE O CUSTO.", "PROBUILT")
bs["I4"].comment = Comment("Usado apenas quando as colunas E–H (decomposição) estão vazias.", "PROBUILT")

wb.save(out); print(len(rows), "artigos")
