import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
OBRA = {"A": "Ambas", "N": "Construção nova", "R": "Remodelação"}
NAT = {"MQ": "Omissão no MQ", "P": "Omissão de projeto", "C": "Compatibilização"}

rows, esps = [], []
code = esp = None; n = 0
for line in open(src, encoding="utf-8"):
    line = line.strip()
    if not line: continue
    if line.startswith("#"):
        code, esp = line[1:].split("|"); n = 0; esps.append(esp); continue
    g, t, o, w, grav, nat = line.split("|")
    n += 1
    rows.append([f"{code}-{n:02d}", esp, g, t, OBRA[o], w, NAT[nat], grav])

F = "Arial"
thin = Side(style="thin", color="BFBFBF")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
INPUT = PatternFill("solid", fgColor="FFF2CC")
BODY = Font(name=F, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()

# ---------- Instruções ----------
ins = wb.active; ins.title = "Instruções"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 2
ins.column_dimensions["B"].width = 26
ins.column_dimensions["C"].width = 95
ins["B1"] = "Checklist de Omissões e Compatibilização de Projetos"
ins["B1"].font = Font(name=F, bold=True, size=16, color="1F3864")
ins["B2"] = "Moradias e remodelações · v1.0 · Ponto de partida para validação e enriquecimento pelo utilizador"
ins["B2"].font = Font(name=F, italic=True, size=10, color="595959")

ins["B3"] = "DADOS DA OBRA"; ins["B3"].font = Font(name=F, bold=True, size=11, color="1F3864")
campos = ["Obra", "Tipo de obra", "Revisão do projeto analisada", "Data da análise", "Analisado por"]
for i, c in enumerate(campos):
    r = 4 + i
    ins.cell(r, 2, c).font = Font(name=F, bold=True, size=10)
    cell = ins.cell(r, 3); cell.fill = INPUT; cell.border = B; cell.font = BODY
ins["C7"].number_format = "dd/mm/yyyy"
dv_obra = DataValidation(type="list", formula1='"Construção nova,Remodelação"', allow_blank=True)
ins.add_data_validation(dv_obra); dv_obra.add("C5")
ins["D5"] = "← escolha da lista: filtra automaticamente os pontos aplicáveis (coluna I da Checklist)"
ins["D5"].font = Font(name=F, italic=True, size=9, color="595959")

r = 10
def sec(title):
    global r
    ins.cell(r, 2, title).font = Font(name=F, bold=True, size=11, color="1F3864"); r += 1
def line(a, b, fill=None):
    global r
    ca = ins.cell(r, 2, a); ca.font = Font(name=F, bold=True, size=10); ca.alignment = WRAP
    cb = ins.cell(r, 3, b); cb.font = BODY; cb.alignment = WRAP
    if fill: ca.fill = fill
    r += 1

sec("COMO USAR (passo a passo)")
for k, v in [
    ("1. Dados da obra", "Preencha as células amarelas acima. O «Tipo de obra» marca como «Não» os pontos que só se aplicam ao outro tipo."),
    ("2. Filtrar", "Na Checklist, filtre a coluna I = «Sim». Trabalhe por especialidade (coluna B)."),
    ("3. Verificar", "Para cada ponto, consulte a peça indicada em «Onde verificar» e escolha o Estado (coluna J)."),
    ("4. Rastrear", "Para cada «Omisso» ou «A esclarecer», preencha o capítulo/artigo do MQ (K), o desenho/peça (L) e a ação proposta (M)."),
    ("5. Valorizar", "Opcional: estime o valor da omissão (N) para medir o impacto no orçamento."),
    ("6. Resumo", "A folha «Resumo» calcula o progresso e as omissões por especialidade e por natureza."),
]: line(k, v)
r += 1
sec("CÉLULAS A PREENCHER")
line("Fundo amarelo", "Únicas células de introdução de dados (Dados da obra e colunas J a N da Checklist). As restantes colunas são a base da checklist ou fórmulas.", INPUT)
r += 1
sec("ESTADOS (coluna J)")
for k, v, c in [("OK", "Verificado e completo no projeto e no MQ.", "C6EFCE"),
                ("Omisso", "Falta no MQ, falta definição no projeto ou há incompatibilidade confirmada.", "FFC7CE"),
                ("A esclarecer", "Dúvida a colocar ao projetista (pedido de esclarecimento).", "FFEB9C"),
                ("N/A", "Não se aplica a esta obra em concreto.", "E7E6E6"),
                ("(vazio)", "Ainda por verificar.", None)]:
    line(k, v, PatternFill("solid", fgColor=c) if c else None)
r += 1
sec("NATUREZA (coluna G)")
for k, v in [("Omissão no MQ", "O trabalho é necessário mas pode não ter artigo no mapa de quantidades."),
             ("Omissão de projeto", "Falta uma definição técnica, peça desenhada ou estudo."),
             ("Compatibilização", "Interface entre especialidades: verificar que as peças dizem o mesmo e não colidem.")]:
    line(k, v)
r += 1
sec("GRAVIDADE TÍPICA (coluna H)")
for k, v in [("Alta", "Impacto relevante no custo, no prazo ou no desempenho/segurança, ou retrabalho difícil."),
             ("Média", "Impacto moderado, resolúvel em obra com custo controlado."),
             ("Baixa", "Impacto reduzido ou facilmente absorvido.")]:
    line(k, v)
ins.cell(r, 3, "A gravidade é indicativa: ajuste-a caso a caso no campo de observações.").font = Font(name=F, italic=True, size=9, color="595959")
r += 2
sec("EXEMPLO DE LINHA PREENCHIDA")
ex_h = ["ID", "Ponto de verificação", "Estado", "Capítulo/Artigo MQ", "Desenho/Peça", "Observações / Ação proposta", "Valor estimado (€)"]
ex_v = ["ARQ-13", "Zonas de duche, bases de banheira e pavimentos de casas de banho", "Omisso", "Cap. 7 – Revestimentos (sem artigo)",
        "ARQ-05 Rev.B – Planta piso 1", "Sem artigo de impermeabilização sob duches à face; pedir artigo ao projetista e confirmar rebaixo na EST-03.", 450]
for i, (h, v) in enumerate(zip(ex_h, ex_v)):
    ins.cell(r + i, 2, h).font = Font(name=F, bold=True, size=10)
    c = ins.cell(r + i, 3, v); c.font = Font(name=F, size=10, color="0000FF"); c.alignment = WRAP
ins.cell(r + 6, 3).number_format = '#,##0 "€"'
ins.cell(r + 6, 3).alignment = Alignment(horizontal="left")
r += 8
ins.cell(r, 2, "Nota").font = Font(name=F, bold=True, size=9)
c = ins.cell(r, 3, "Checklist genérica elaborada com base em boas práticas de construção em Portugal. Não substitui a verificação regulamentar específica de cada projeto. Deve ser enriquecida com a experiência do utilizador (acrescentar linhas no fim da Checklist, mantendo o formato do ID).")
c.font = Font(name=F, italic=True, size=9, color="595959"); c.alignment = WRAP

# ---------- Checklist ----------
ck = wb.create_sheet("Checklist")
ck["A1"] = "CHECKLIST DE OMISSÕES E COMPATIBILIZAÇÃO"
ck["A1"].font = Font(name=F, bold=True, size=14, color="1F3864")
ck["A2"] = '="Obra: "&IF(\'Instruções\'!$C$4="","(preencher em Instruções)",\'Instruções\'!$C$4)&"   ·   Tipo: "&IF(\'Instruções\'!$C$5="","(todos os pontos)",\'Instruções\'!$C$5)'
ck["A2"].font = Font(name=F, italic=True, size=10, color="595959")
heads = ["ID", "Especialidade", "Grupo", "Ponto de verificação", "Aplica-se a", "Onde verificar", "Natureza", "Gravidade típica",
         "Aplicável nesta obra", "Estado", "Capítulo / Artigo do MQ", "Desenho / Peça de referência", "Observações / Ação proposta", "Valor estimado (€)"]
widths = [9, 20, 17, 60, 15, 24, 18, 11, 11, 13, 22, 24, 45, 13]
HR = 3
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = ck.cell(HR, i, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = B
    ck.column_dimensions[get_column_letter(i)].width = w
for i in range(10, 15):
    ck.cell(HR, i).fill = PatternFill("solid", fgColor="BF8F00")
ck.row_dimensions[HR].height = 32

first = HR + 1
for k, row in enumerate(rows):
    rr = first + k
    for j, v in enumerate(row, 1):
        c = ck.cell(rr, j, v); c.font = BODY; c.border = B
        c.alignment = CENTER if j in (1, 5, 8) else WRAP
    ck.cell(rr, 9, f'=IF(OR(E{rr}="Ambas",\'Instruções\'!$C$5="",E{rr}=\'Instruções\'!$C$5),"Sim","Não")')
    ck.cell(rr, 9).font = BODY; ck.cell(rr, 9).alignment = CENTER; ck.cell(rr, 9).border = B
    for j in range(10, 15):
        c = ck.cell(rr, j); c.fill = INPUT; c.border = B; c.font = BODY; c.alignment = CENTER if j == 10 else WRAP
    ck.cell(rr, 14).number_format = '#,##0 "€";-#,##0 "€";-'
last = first + len(rows) - 1
# linhas extra em branco para o utilizador acrescentar pontos
EXTRA = 30
for rr in range(last + 1, last + 1 + EXTRA):
    for j in range(1, 15):
        c = ck.cell(rr, j); c.border = B; c.font = BODY; c.alignment = WRAP
        if j >= 10 or j <= 8: c.fill = INPUT if j >= 10 else PatternFill("solid", fgColor="F2F2F2")
    ck.cell(rr, 9, f'=IF(D{rr}="","",IF(OR(E{rr}="Ambas",E{rr}="",\'Instruções\'!$C$5="",E{rr}=\'Instruções\'!$C$5),"Sim","Não"))')
    ck.cell(rr, 9).alignment = CENTER
    ck.cell(rr, 14).number_format = '#,##0 "€";-#,##0 "€";-'
end = last + EXTRA

dv_est = DataValidation(type="list", formula1='"OK,Omisso,A esclarecer,N/A"', allow_blank=True)
dv_grav = DataValidation(type="list", formula1='"Alta,Média,Baixa"', allow_blank=True)
dv_aplica = DataValidation(type="list", formula1='"Ambas,Construção nova,Remodelação"', allow_blank=True)
dv_nat = DataValidation(type="list", formula1='"Omissão no MQ,Omissão de projeto,Compatibilização"', allow_blank=True)
dv_esp = DataValidation(type="list", formula1=f"=Resumo!$A$5:$A${4+len(esps)}", allow_blank=True)
for dv in (dv_est, dv_grav, dv_aplica, dv_nat, dv_esp): ck.add_data_validation(dv)
dv_est.add(f"J{first}:J{end}")
dv_grav.add(f"H{last+1}:H{end}"); dv_aplica.add(f"E{last+1}:E{end}"); dv_nat.add(f"G{last+1}:G{end}"); dv_esp.add(f"B{last+1}:B{end}")

rng = f"A{first}:N{end}"
ck.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{first}="Não"'], font=Font(color="A6A6A6"), fill=PatternFill("solid", fgColor="F2F2F2"), stopIfTrue=True))
for val, col in [("OK", "C6EFCE"), ("Omisso", "FFC7CE"), ("A esclarecer", "FFEB9C"), ("N/A", "E7E6E6")]:
    ck.conditional_formatting.add(f"J{first}:J{end}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=col)))
ck.conditional_formatting.add(f"H{first}:H{end}", CellIsRule(operator="equal", formula=['"Alta"'], font=Font(color="C00000", bold=True)))
ck.freeze_panes = f"E{first}"
ck.auto_filter.ref = f"A{HR}:N{end}"
ck.print_title_rows = f"{HR}:{HR}"
ck.page_setup.orientation = "landscape"; ck.page_setup.fitToWidth = 1; ck.page_setup.fitToHeight = 0
ck.sheet_properties.pageSetUpPr.fitToPage = True

# ---------- Resumo ----------
rs = wb.create_sheet("Resumo")
rs.sheet_view.showGridLines = False
rs["A1"] = "RESUMO DA VERIFICAÇÃO"; rs["A1"].font = Font(name=F, bold=True, size=14, color="1F3864")
rs["A2"] = ck["A2"].value; rs["A2"].font = Font(name=F, italic=True, size=10, color="595959")
rh = ["Especialidade", "Pontos aplicáveis", "OK", "Omisso", "A esclarecer", "N/A", "Por verificar", "% verificado", "Omissões gravidade Alta", "Valor estimado (€)"]
rw = [30, 12, 9, 9, 11, 9, 11, 11, 13, 15]
for i, (h, w) in enumerate(zip(rh, rw), 1):
    c = rs.cell(4, i, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.border = B
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rs.column_dimensions[get_column_letter(i)].width = w
rs.row_dimensions[4].height = 32
CB = f"Checklist!$B${first}:$B${end}"; CI = f"Checklist!$I${first}:$I${end}"
CJ = f"Checklist!$J${first}:$J${end}"; CH = f"Checklist!$H${first}:$H${end}"
CN = f"Checklist!$N${first}:$N${end}"; CG = f"Checklist!$G${first}:$G${end}"
for k, e in enumerate(esps):
    rr = 5 + k
    rs.cell(rr, 1, e)
    rs.cell(rr, 2, f'=COUNTIFS({CB},$A{rr},{CI},"Sim")')
    for col, st in zip("CDEF", ["OK", "Omisso", "A esclarecer", "N/A"]):
        rs[f"{col}{rr}"] = f'=COUNTIFS({CB},$A{rr},{CI},"Sim",{CJ},"{st}")'
    rs[f"G{rr}"] = f"=B{rr}-SUM(C{rr}:F{rr})"
    rs[f"H{rr}"] = f'=IF(B{rr}=0,"-",(B{rr}-G{rr})/B{rr})'
    rs[f"I{rr}"] = f'=COUNTIFS({CB},$A{rr},{CI},"Sim",{CJ},"Omisso",{CH},"Alta")'
    rs[f"J{rr}"] = f'=SUMIFS({CN},{CB},$A{rr},{CI},"Sim")'
tr = 5 + len(esps)
rs.cell(tr, 1, "TOTAL")
for col in "BCDEFGIJ":
    rs[f"{col}{tr}"] = f"=SUM({col}5:{col}{tr-1})"
rs[f"H{tr}"] = f'=IF(B{tr}=0,"-",(B{tr}-G{tr})/B{tr})'
for rr in range(5, tr + 1):
    for j in range(1, 11):
        c = rs.cell(rr, j); c.border = B
        c.font = Font(name=F, size=10, bold=(rr == tr))
        if j > 1: c.alignment = Alignment(horizontal="center")
        if rr == tr: c.fill = PatternFill("solid", fgColor="D9E1F2")
    rs.cell(rr, 8).number_format = "0%"
    rs.cell(rr, 10).number_format = '#,##0 "€";-#,##0 "€";-'
rs.conditional_formatting.add(f"D5:D{tr-1}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="C00000", bold=True)))
rs.conditional_formatting.add(f"I5:I{tr-1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE")))

r2 = tr + 3
rs.cell(r2 - 1, 1, "POR NATUREZA").font = Font(name=F, bold=True, size=11, color="1F3864")
for i, h in enumerate(["Natureza", "Pontos aplicáveis", "OK", "Omisso", "A esclarecer", "N/A", "Por verificar"], 1):
    c = rs.cell(r2, i, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.border = B; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for k, nname in enumerate(["Omissão no MQ", "Omissão de projeto", "Compatibilização"]):
    rr = r2 + 1 + k
    rs.cell(rr, 1, nname)
    rs[f"B{rr}"] = f'=COUNTIFS({CG},$A{rr},{CI},"Sim")'
    for col, st in zip("CDEF", ["OK", "Omisso", "A esclarecer", "N/A"]):
        rs[f"{col}{rr}"] = f'=COUNTIFS({CG},$A{rr},{CI},"Sim",{CJ},"{st}")'
    rs[f"G{rr}"] = f"=B{rr}-SUM(C{rr}:F{rr})"
    for j in range(1, 8):
        c = rs.cell(rr, j); c.border = B; c.font = BODY
        if j > 1: c.alignment = Alignment(horizontal="center")
rs.cell(r2 + 5, 1, "Fórmulas: contagens com COUNTIFS sobre a folha Checklist, considerando apenas pontos com «Aplicável nesta obra» = Sim. «Por verificar» = aplicáveis − (OK + Omisso + A esclarecer + N/A).").font = Font(name=F, italic=True, size=9, color="595959")

wb.move_sheet("Resumo", offset=-1)
wb.active = 0
wb.save(out)
print(len(rows), "pontos;", len(esps), "especialidades")
