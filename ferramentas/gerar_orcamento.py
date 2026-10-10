"""Gerador de orçamento DRAFT PROBUILT.

Lê um mapa de quantidades (Excel) do projetista, mantém a estrutura de capítulos,
propõe preços a partir da BASE_PRECOS_PROBUILT e escreve o resultado na folha de
orçamento PROBUILT (_Produção_Folha_de_orçamento.xlsx), com rastreabilidade.

Uso:
  python3 gerar_orcamento.py --mq mapa.xlsx --base BASE_PRECOS.xlsx --modelo folha.xlsx \
      --saida ORC_vDRAFT.xlsx [--folha "Nome da folha"] [--obra "Nome"] [--cliente "Nome"]
"""
import argparse
import math
import re
import unicodedata
from collections import Counter
from datetime import date

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LIMIAR_PROPOSTO = 0.60
LIMIAR_CONFIRMAR = 0.45

STOP = set("""a o os as de da do das dos e em no na nos nas com sem por para ao aos um uma uns umas que se
ou incl incluindo inclui incluido fornecimento fornecer execucao executar aplicacao aplicar todos todas
trabalhos trabalho necessarios necessario conforme pecas desenhadas desenhado projeto tipo equivalente
etc ate sobre sob entre cada outro outros outras outra este esta estes estas seu sua seus suas pelo pela
respetivo respetiva respetivos devidamente acordo forma demais materiais material acessorios verba""".split())

UNIDADES = {"m2": "m2", "m²": "m2", "m3": "m3", "m³": "m3", "ml": "m", "m.l.": "m", "m": "m", "mts": "m",
            "un": "un", "un.": "un", "und": "un", "uni": "un", "unid": "un", "u": "un", "pç": "un", "pc": "un",
            "vg": "vg", "vb": "vg", "verba": "vg", "cj": "vg", "conj": "vg", "gl": "vg", "kg": "kg", "kw": "kw", "h": "h"}

# Elementos construtivos (tokens já truncados a 6 letras): artigos de elementos diferentes não se correspondem
ELEMENTOS = {"parede", "pavime", "teto", "tecto", "cobert", "fachad", "escada", "porta", "janela", "caixil",
             "laje", "pilar", "pilare", "viga", "sapata", "muro"}

# Trabalhos de demolição/remoção não se correspondem com trabalhos de execução (e vice-versa)
DEMOLICAO = {"demoli", "remoca", "remove", "desmon", "arranq", "levant"}
FORNECE_DO = r"a fornecer (e escolher )?pelo dono|fornecid[oa]s? pelo dono|fornecimento do dono|a fornecer e escolher pelo"

# Rastreio automático de omissões: (ID checklist, ponto, regex que deve existir no MQ, regex de ativação opcional)
OMISSOES = [
    ("GER-01", "Estaleiro (montagem, manutenção, desmontagem, proteções)", r"estaleiro|tapume", None),
    ("GER-03", "Andaimes / plataformas de trabalho", r"andaime|plataforma", r"fachada|cobertura|exterior"),
    ("GER-06", "Plano de Segurança e Saúde / proteções coletivas", r"seguranca e saude|\bpss\b|protecoes coletivas", None),
    ("GER-07", "Gestão de RCD / transporte a vazadouro", r"residuo|\brcd\b|vazadouro|entulho", None),
    ("GER-09", "Limpeza final de obra", r"limpeza", None),
    ("GER-10", "Telas finais / compilação técnica", r"telas finais|compilacao tecnica", None),
    ("DEM-08", "Picagem / reparação de suportes após demolições", r"picagem|regulariza|reparacao de", r"demoli|remoc"),
    ("ARQ-08", "Rodapés", r"rodape", r"pavimento"),
    ("ARQ-13", "Impermeabilização de zonas de duche / I.S.", r"impermeabiliz", r"duche|banho|sanita|\bi\.?s\.?\b|\bwc\b"),
    ("ARQ-30", "Portas interiores (aros, guarnições, ferragens)", r"\bporta", r"remocao de portas|demolicao de portas"),
    ("ARQ-38", "Loiças sanitárias e torneiras", r"loica|sanita|lavatorio|torneira|base de duche", r"duche|banho|sanitari"),
    ("ARQ-36", "Pintura interior de paredes e tetos", r"pintura", None),
    ("AAG-11", "Ensaios de estanquidade (águas/esgotos)", r"ensaio", r"abastecimento|drenagem|esgoto|canaliza|pichelaria"),
    ("ELE-15", "Certificação / inspeção da instalação elétrica", r"certifica|inspec", r"eletric|electric"),
    ("ITE-05", "Certificação ITED", r"ited|certifica", r"\bited\b|telecomunica"),
    ("LIC-01", "Taxas e ligações às redes", r"taxa|ramal|ligacao a rede", r"construcao nova|moradia nova|ramal"),
]

fill = lambda c: PatternFill("solid", fgColor=c)
VERDE, AMARELO, VERMELHO, LARANJA, AZUL = fill("C6EFCE"), fill("FFEB9C"), fill("FFC7CE"), fill("F8CBAD"), fill("DDEBF7")
CAP_FILL = {1: fill("D9E1F2"), 2: fill("EDF2F9"), 3: fill("F5F8FC")}
FONTE = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)


def norm(txt):
    txt = str(txt or "").replace("Ø", " d").replace("ø", " d")   # o «Ø» perde-se na normalização ASCII
    txt = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", txt)


def tokens(txt):
    out = []
    for t in re.findall(r"[a-z]+\d+[a-z\d]*|[a-z]+|\d+(?:[.,]\d+)?", norm(txt)):
        if t in STOP or (len(t) < 3 and not t[0].isdigit()) or (t.isdigit() and len(t) == 1):
            continue   # códigos como «ba13», «c25» ficam inteiros; números de 1 dígito são ambíguos (8 mm vs 8 cm)
        if not t[0].isdigit() and len(t) > 4 and t.endswith("s"):
            t = t[:-1]                      # plural simples: paredes → parede, tetos → teto
        out.append(t[:6] if t.isalpha() else t)
    return out


def diametros(txt):
    return set(re.findall(r"(?:\bd|ø|diametro)\s?(\d{2,3})\b", norm(txt)))


def unidade(u):
    u = norm(u).strip().replace(" ", "")
    return UNIDADES.get(u, u)


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            s = s[1:]
            if re.fullmatch(r"[\d.,+\-*/() ]+", s):
                try:
                    return float(eval(s.replace(",", ".")))  # só aritmética simples, validada pelo regex
                except Exception:
                    return None
            return None
        s = s.replace("€", "").replace(" ", "")
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?", s):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None
    return None


# ------------------------------------------------------------------ leitura do MQ
def encontrar_cabecalho(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        cols = {}
        for c in row:
            t = norm(c.value)
            if not t:
                continue
            if "desc" in t and "desc" not in cols:
                cols["desc"] = c.column
            elif re.fullmatch(r"un\.?|unid\.?|unidade", t.strip()):
                cols["un"] = c.column
            elif re.match(r"quant|qtd|qt\.", t.strip()):
                cols["qtd"] = c.column
            elif re.fullmatch(r"art.*|ref\.?|item|cod.*|n\.?o?", t.strip()):
                cols["cod"] = c.column
        if {"desc", "un", "qtd"} <= cols.keys():
            cols.setdefault("cod", cols["desc"] - 1)
            return row[0].row, cols
    return None, None


def nivel(cod):
    partes = [p for p in re.split(r"[.\s]", str(cod).strip()) if p]
    return max(1, min(len(partes), 3))


def ler_mq(caminho, folha=None):
    wb_v = load_workbook(caminho, data_only=True)
    wb_f = load_workbook(caminho, data_only=False)
    candidatas = [folha] if folha else wb_v.sheetnames
    melhor = None
    for nome in candidatas:
        hr, cols = encontrar_cabecalho(wb_v[nome])
        if not hr:
            continue
        itens = _extrair(wb_v[nome], wb_f[nome], hr, cols)
        n_art = sum(1 for i in itens if i["tipo"] == "artigo")
        if not melhor or n_art > melhor[2]:
            melhor = (nome, itens, n_art)
    if not melhor:
        raise SystemExit("Não encontrei cabeçalho (Descrição / Un / Quant.) em nenhuma folha do MQ.")
    return melhor[0], melhor[1]


def _extrair(ws, wsf, hr, cols):
    linhas = []
    for r in range(hr + 1, ws.max_row + 1):
        cod = ws.cell(r, cols["cod"]).value
        desc = ws.cell(r, cols["desc"]).value
        un = ws.cell(r, cols["un"]).value
        q = ws.cell(r, cols["qtd"]).value
        if q is None:
            q = wsf.cell(r, cols["qtd"]).value
        cod = str(cod).strip() if cod not in (None, "") else ""
        desc = str(desc).strip() if desc not in (None, "") else ""
        if not cod and not desc:
            continue
        if not desc and len(cod) > 6:          # ex.: "A — DEMOLIÇÕES" na coluna do código
            m = re.match(r"^\s*([A-Z0-9][\w.]*)\s*[—–-]\s*(.+)$", cod)
            cod, desc = (m.group(1), m.group(2)) if m else ("", cod)
        nd = norm(desc)
        if re.match(r"^(sub)?total|^iva\b|^total geral", nd) or "fundo laranja" in nd:
            continue
        linhas.append(dict(linha=r, cod=cod, desc=desc, un=str(un).strip() if un not in (None, "") else "", qtd=num(q)))

    itens, titulo, ult_cod = [], None, None
    sufixo = Counter()
    for i, l in enumerate(linhas):
        e_artigo = bool(l["un"]) or (l["qtd"] is not None and l["qtd"] != 0 and not l["cod"])
        if e_artigo:
            cod = l["cod"]
            desc = l["desc"]
            if not cod and titulo:
                sufixo[titulo["cod"]] += 1
                cod = f'{titulo["cod"]}.{chr(96 + sufixo[titulo["cod"]])}' if sufixo[titulo["cod"]] > 1 or _mais_filhos(linhas, i) else titulo["cod"]
                desc = f'{titulo["desc"]} – {desc}'
            elif not cod:
                sufixo[ult_cod] += 1         # linha sem código a seguir a um artigo: herda o código + letra
                cod = f"{ult_cod}.{chr(96 + sufixo[ult_cod])}" if ult_cod else "s/ ref."
            if l["cod"]:
                ult_cod = l["cod"]
            elif titulo:
                ult_cod = titulo["cod"]
            itens.append(dict(tipo="artigo", cod=cod, desc=desc, un=l["un"], qtd=l["qtd"] or 0, linha=l["linha"],
                              titulo=titulo["desc"] if titulo and desc.startswith(titulo["desc"]) else ""))
            continue
        if l["cod"]:
            prox = next((x for x in linhas[i + 1:]), None)
            if prox and not prox["cod"] and prox["un"]:
                titulo = dict(cod=l["cod"], desc=l["desc"])   # título de artigo; descrição vem na linha seguinte
                continue
            titulo = None
            itens.append(dict(tipo="capitulo", cod=l["cod"].rstrip("."), desc=l["desc"], nivel=nivel(l["cod"]), linha=l["linha"]))
        else:
            itens.append(dict(tipo="nota", desc=l["desc"], linha=l["linha"]))
    return itens


def _mais_filhos(linhas, i):
    seguinte = linhas[i + 1] if i + 1 < len(linhas) else None
    return bool(seguinte and not seguinte["cod"] and seguinte["un"])


# ------------------------------------------------------------------ base de preços
def ler_base(caminho):
    ws = load_workbook(caminho, data_only=True)["Base"]
    wsf = load_workbook(caminho)["Base"]
    base = []
    for r in range(5, ws.max_row + 1):
        cod = ws.cell(r, 1).value
        if not cod:
            continue
        g = lambda c: num(wsf.cell(r, c).value) or 0.0
        base.append(dict(cod=cod, desc=ws.cell(r, 2).value or "", esp=ws.cell(r, 3).value, un=unidade(ws.cell(r, 4).value),
                         mo=g(5), mat=g(6), eq=g(7), sub=g(8), glob=g(9), fonte=ws.cell(r, 15).value, obra=ws.cell(r, 16).value,
                         fiab=ws.cell(r, 18).value, margem=num(wsf.cell(r, 12).value)))
    df = Counter()
    for b in base:
        b["tok"] = set(tokens(b["desc"]))
        df.update(b["tok"])
    n = len(base)
    idf = {t: math.log((n + 1) / (c + 0.5)) for t, c in df.items()}
    return base, idf


def melhor_preco(item, base, idf):
    tok_a = set(tokens(item["desc"]))
    inicio_a = tokens(item.get("titulo") or item["desc"])[:4]
    un_a = unidade(item["un"])
    res = []
    for b in base:
        comuns = tok_a & b["tok"]
        if not comuns:
            continue
        w = lambda s: sum(idf.get(t, 3.0) for t in s)
        cobertura_base = w(comuns) / max(w(b["tok"]), 1e-9)
        cobertura_art = w(comuns) / max(w(tok_a), 1e-9)
        score = 0.75 * cobertura_base + 0.25 * min(1.0, cobertura_art * 2)
        if un_a and b["un"] and un_a != b["un"]:
            score *= 0.5
        dia_a, dia_b = diametros(item["desc"]), diametros(b["desc"])
        if dia_a and dia_b and not (dia_a & dia_b):   # diâmetros diferentes (ex.: Ø90 vs D110)
            score *= 0.6
        if bool(set(inicio_a[:3]) & DEMOLICAO) != bool(set(tokens(b["desc"])[:3]) & DEMOLICAO):   # natureza lê-se no verbo inicial
            score *= 0.5
        tit = set(tokens(re.sub(r"^\s*[A-Za-z]{1,3}\.?\d+\s*\|", "", item.get("titulo", ""))))
        if tit and not (tit & b["tok"]):          # o título do artigo (ex.: «REGULARIZAÇÃO») tem de estar na base
            score *= 0.6
        if inicio_a and inicio_a[0] not in b["tok"] and inicio_a[0] not in ELEMENTOS:
            score *= 0.75                         # o trabalho principal do artigo (1.ª palavra útil) não está na base
        el_a, el_b = tok_a & ELEMENTOS, b["tok"] & ELEMENTOS
        if el_a and el_b and not (el_a & el_b):   # elemento construtivo diferente (ex.: parede vs pavimento)
            score *= 0.6
        res.append((round(score, 3), b, w(comuns)))
    res.sort(key=lambda x: (-x[0], -x[2], x[1]["fiab"] or "Z"))   # empate: ganha quem cobre mais do artigo
    res = [(sc, b) for sc, b, _ in res]
    return res[:3]


def rastrear_omissoes(itens):
    texto = norm(" ".join(i["desc"] for i in itens if i["tipo"] != "nota"))
    out = []
    for cid, ponto, deve, ativa in OMISSOES:
        if ativa and not re.search(ativa, texto):
            continue
        ok = re.search(deve, texto)
        out.append((cid, ponto, "Encontrado no MQ" if ok else "POSSÍVEL OMISSÃO"))
    return out


# ------------------------------------------------------------------ escrita na folha de orçamento
def escrever(modelo, saida, itens, base, idf, meta):
    wb = load_workbook(modelo)
    ws = wb["Orçamento"]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 8:
            ws.unmerge_cells(str(rng))
    for r in range(8, ws.max_row + 1):
        for c in range(1, 22):
            ws.cell(r, c).value = None
    ws["C2"], ws["C3"], ws["C4"], ws["C5"] = meta.get("cliente", ""), meta.get("obra", ""), meta.get("morada", ""), date.today().strftime("%d/%m/%Y")
    ws["B1"] = f'Orçamento nº {meta.get("numero", "DRAFT")}'   # B1:D1 está unida no modelo

    extra = ["Cód. base", "Designação na base", "Score", "Estado", "Fiab.", "Fonte do preço", "Linha MQ"]
    for i, h in enumerate(extra):
        c = ws.cell(7, 15 + i, h)
        c.font = Font(name=FONTE, bold=True, size=9)
        c.fill = AZUL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for col, w in zip("OPQRSTU", [9, 40, 7, 13, 6, 30, 8]):
        ws.column_dimensions[col].width = w

    r = 8
    rast, caps = [], []
    for it in itens:
        if it["tipo"] == "nota":
            continue
        if it["tipo"] == "capitulo":
            ws.cell(r, 1, it["nivel"]); ws.cell(r, 2, it["cod"]); ws.cell(r, 3, it["desc"])
            for c in range(1, 15):
                ws.cell(r, c).fill = CAP_FILL[it["nivel"]]
                ws.cell(r, c).font = Font(name=FONTE, bold=True, size=10 if it["nivel"] == 1 else 9)
            if it["nivel"] == 1:
                caps.append([it["cod"], it["desc"], r, None])
        else:
            ws.cell(r, 2, it["cod"]); ws.cell(r, 3, it["desc"]); ws.cell(r, 4, it["un"]); ws.cell(r, 5, it["qtd"])
            ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
            cand = melhor_preco(it, base, idf)
            score, b = cand[0] if cand else (0, None)
            estado = "Proposto" if score >= LIMIAR_PROPOSTO else "A confirmar" if score >= LIMIAR_CONFIRMAR else "Sem preço"
            material_do = bool(re.search(FORNECE_DO, norm(it["desc"])))
            if material_do and estado == "Proposto":
                estado = "A confirmar"           # preço da base inclui material; aqui o material é do dono de obra
            if b and estado != "Sem preço":
                if b["mo"] or b["mat"] or b["eq"] or b["sub"]:
                    ws.cell(r, 10, b["mo"] or None); ws.cell(r, 11, b["mat"] or None); ws.cell(r, 12, b["eq"] or None); ws.cell(r, 13, b["sub"] or None)
                else:
                    ws.cell(r, 13, b["glob"])
                    ws.cell(r, 13).comment = Comment("Preço global da base (sem decomposição) colocado em Subempreitada.", "Gerador PROBUILT")
                for c in (10, 11, 12, 13):
                    ws.cell(r, c).font = Font(name=FONTE, size=10, color="0000FF")
            vals = [b["cod"] if b else "", b["desc"] if b else "", score, estado, (b["fiab"] if b else ""), (b["fonte"] if b else ""), it["linha"]]
            for i, v in enumerate(vals):
                c = ws.cell(r, 15 + i, v)
                c.font = Font(name=FONTE, size=8)
                c.alignment = Alignment(wrap_text=i in (1, 5), vertical="top")
            ws.cell(r, 18).fill = {"Proposto": VERDE, "A confirmar": AMARELO, "Sem preço": VERMELHO}[estado]
            if material_do:
                ws.cell(r, 18).comment = Comment("Material fornecido pelo dono de obra: o preço deve ser só de mão de obra/aplicação.", "Gerador PROBUILT")
                ws.cell(r, 3).fill = LARANJA
            alt = " | ".join(f'{x[1]["cod"]} ({x[0]:.2f})' for x in cand[1:])
            rast.append([it["linha"], it["cod"], it["desc"], it["un"], it["qtd"], vals[0], vals[1], score, estado, vals[4], vals[5], alt])
        ws.cell(r, 6, f"=ROUND((1+I{r})*N{r},2)")
        ws.cell(r, 7, f"=ROUND(F{r}*E{r},2)")
        ws.cell(r, 9, f"=+IF(N{r}<>0,$I$6,0)")
        ws.cell(r, 14, f'=IF(E{r}="Incluido",0,SUM(J{r}*IF(J$6=0,1,(1+J$6))+K{r}*IF(K$6=0,1,(1+$K$6))+L{r}*IF(L$6=0,1,(1+$L$6))+M{r}*IF(M$6=0,1,(1+$M$6))))')
        for c, fmt in ((5, "#,##0.00"), (6, '#,##0.00\\ "€"'), (7, '#,##0.00\\ "€"'), (9, "0%"), (14, '#,##0.00\\ "€"')):
            ws.cell(r, c).number_format = fmt
        r += 1
    ult = r - 1
    for i, cap in enumerate(caps):
        cap[3] = (caps[i + 1][2] - 1) if i + 1 < len(caps) else ult

    tot = ult + 2
    ws.cell(tot, 2, "TOTAL").font = Font(name=FONTE, bold=True)
    ws.cell(tot, 7, f"=SUM(G8:G{ult})")
    for col in "JKLM":
        ws[f"{col}{tot}"] = f"=SUMPRODUCT($E$8:$E${ult},{col}8:{col}{ult})"
    ws[f"N{tot}"] = f"=SUMPRODUCT($E$8:$E${ult},N8:N{ult})"   # custo total c/ encargos ("Seco") – em falta no modelo original
    for col in "GJKLMN":
        ws[f"{col}{tot}"].number_format = '#,##0.00\\ "€"'
        ws[f"{col}{tot}"].font = Font(name=FONTE, bold=True)
    ws["J1"] = f'=IF(COUNTIF(E8:E{ult},">0")-COUNTIF(F8:F{ult},">0")<>0,COUNTIF(E8:E{ult},">0")-COUNTIF(F8:F{ult},">0")&" preços",0)'
    ws["J2"] = f"=+G{tot}"
    ws["J3"] = f"=N{tot}"
    ws.print_area = f"B8:G{tot}"
    ws.freeze_panes = "D8"

    # ---- Resumo por capítulo
    rs = wb.create_sheet("Resumo")
    rs["A1"] = f'RESUMO – {meta.get("obra", "")}'; rs["A1"].font = Font(name=FONTE, bold=True, size=14, color="1F3864")
    h = ["Cap.", "Capítulo", "Custo c/ encargos (€)", "Preço de venda (€)", "Margem (€)", "Margem s/ custo", "Artigos", "Sem preço"]
    for i, t in enumerate(h, 1):
        c = rs.cell(3, i, t); c.font = Font(name=FONTE, bold=True, color="FFFFFF"); c.fill = fill("1F3864"); c.alignment = Alignment(wrap_text=True, horizontal="center")
    for k, (cod, desc, a, b_) in enumerate(caps):
        rr = 4 + k
        rs.cell(rr, 1, cod); rs.cell(rr, 2, desc)
        rs.cell(rr, 3, f"=SUMPRODUCT('Orçamento'!E{a}:E{b_},'Orçamento'!N{a}:N{b_})")
        rs.cell(rr, 4, f"=SUM('Orçamento'!G{a}:G{b_})")
        rs.cell(rr, 5, f"=D{rr}-C{rr}")
        rs.cell(rr, 6, f'=IF(C{rr}=0,"-",D{rr}/C{rr}-1)')
        rs.cell(rr, 7, f"=COUNTIF('Orçamento'!D{a}:D{b_},\"?*\")")
        rs.cell(rr, 8, f"=COUNTIF('Orçamento'!R{a}:R{b_},\"Sem preço\")")
    tr = 4 + len(caps)
    rs.cell(tr, 2, "TOTAL")
    for col in "CDEGH":
        rs[f"{col}{tr}"] = f"=SUM({col}4:{col}{tr - 1})"
    rs[f"F{tr}"] = f'=IF(C{tr}=0,"-",D{tr}/C{tr}-1)'
    for rr in range(4, tr + 1):
        for c in range(1, 9):
            cell = rs.cell(rr, c); cell.border = BORDA; cell.font = Font(name=FONTE, size=10, bold=rr == tr)
        for c in (3, 4, 5):
            rs.cell(rr, c).number_format = '#,##0.00 "€"'
        rs.cell(rr, 6).number_format = "0.0%"
    for col, w in zip("ABCDEFGH", [7, 42, 16, 16, 14, 11, 9, 10]):
        rs.column_dimensions[col].width = w
    n_est = Counter(x[8] for x in rast)
    rs.cell(tr + 2, 2, "Estado dos preços").font = Font(name=FONTE, bold=True)
    for k, (e, f) in enumerate([("Proposto", VERDE), ("A confirmar", AMARELO), ("Sem preço", VERMELHO)]):
        rs.cell(tr + 3 + k, 2, e).fill = f
        rs.cell(tr + 3 + k, 3, n_est.get(e, 0))
    rs.cell(tr + 7, 2, f"Limiares: Proposto ≥ {LIMIAR_PROPOSTO:.2f} · A confirmar ≥ {LIMIAR_CONFIRMAR:.2f} · abaixo = sem preço (não preenchido).").font = Font(name=FONTE, italic=True, size=9, color="595959")
    rs.cell(tr + 8, 2, f'MQ de origem: {meta.get("mq", "")} · folha «{meta.get("folha", "")}» · gerado em {date.today():%d/%m/%Y}').font = Font(name=FONTE, italic=True, size=9, color="595959")

    # ---- Rastreabilidade
    rt = wb.create_sheet("Rastreabilidade")
    hh = ["Linha MQ", "Art.º", "Designação (MQ)", "Un", "Quant.", "Cód. base", "Designação na base", "Score", "Estado", "Fiab.", "Fonte", "Alternativas (score)"]
    for i, t in enumerate(hh, 1):
        c = rt.cell(1, i, t); c.font = Font(name=FONTE, bold=True, color="FFFFFF"); c.fill = fill("1F3864"); c.alignment = Alignment(wrap_text=True)
    for k, row in enumerate(rast, 2):
        for i, v in enumerate(row, 1):
            c = rt.cell(k, i, v); c.font = Font(name=FONTE, size=9); c.alignment = Alignment(wrap_text=i in (3, 7, 11, 12), vertical="top")
        rt.cell(k, 9).fill = {"Proposto": VERDE, "A confirmar": AMARELO, "Sem preço": VERMELHO}[row[8]]
    for col, w in zip("ABCDEFGHIJKL", [8, 9, 60, 5, 9, 9, 40, 7, 12, 6, 28, 24]):
        rt.column_dimensions[col].width = w
    rt.freeze_panes = "C2"; rt.auto_filter.ref = f"A1:L{len(rast) + 1}"

    # ---- Omissões (rastreio)
    om = wb.create_sheet("Omissões (rastreio)")
    for i, t in enumerate(["ID checklist", "Ponto de verificação", "Resultado do rastreio", "Decisão (Incluir / OK / N/A)", "Observações"], 1):
        c = om.cell(1, i, t); c.font = Font(name=FONTE, bold=True, color="FFFFFF"); c.fill = fill("1F3864")
    for k, (cid, p, res) in enumerate(rastrear_omissoes(itens), 2):
        om.cell(k, 1, cid); om.cell(k, 2, p); om.cell(k, 3, res).fill = LARANJA if res.startswith("POSS") else VERDE
        om.cell(k, 4).fill = fill("FFF2CC"); om.cell(k, 5).fill = fill("FFF2CC")
    om.cell(k + 2, 2, "Rastreio por palavras-chave: indica onde olhar, não substitui a verificação com a Checklist de Omissões completa.").font = Font(name=FONTE, italic=True, size=9, color="595959")
    for col, w in zip("ABCDE", [11, 50, 22, 24, 40]):
        om.column_dimensions[col].width = w

    wb.save(saida)
    return rast, caps


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mq", required=True); ap.add_argument("--base", required=True); ap.add_argument("--modelo", required=True)
    ap.add_argument("--saida", required=True); ap.add_argument("--folha"); ap.add_argument("--obra", default=""); ap.add_argument("--cliente", default="")
    a = ap.parse_args()
    folha, itens = ler_mq(a.mq, a.folha)
    base, idf = ler_base(a.base)
    rast, caps = escrever(a.modelo, a.saida, itens, base, idf, dict(obra=a.obra, cliente=a.cliente, mq=a.mq.split("/")[-1], folha=folha))
    c = Counter(x[8] for x in rast)
    print(f"Folha: {folha} | capítulos nível 1: {len(caps)} | artigos: {len(rast)} | " + " | ".join(f"{k}: {v}" for k, v in c.items()))


if __name__ == "__main__":
    main()
