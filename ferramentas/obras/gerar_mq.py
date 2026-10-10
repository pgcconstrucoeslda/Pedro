"""Gera o Excel de medição preliminar (Cálculos + MQ) a partir de medicao_apulia.py."""
import re, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from medicao_apulia import CALCULOS, MQ

out = sys.argv[1]
F = "Arial"; thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="1F3864"); HFont = Font(name=F, bold=True, color="FFFFFF", size=10)
CONF = {"A": ("C6EFCE", "Medição do projeto"), "B": ("DDEBF7", "Calculado a partir de dados do projeto"),
        "C": ("FFEB9C", "Estimativa por rácio / pressuposto"), "D": ("FFC7CE", "Verba a medir / a confirmar")}
ref = {cid: i + 4 for i, (cid, *_) in enumerate(CALCULOS)}
TOK = re.compile(r"\b([DEAHSPLV]\d{2})\b")
def conv(expr):
    return "=" + TOK.sub(lambda m: f"Cálculos!$E${ref[m.group(1)]}", expr)

wb = Workbook()
ins = wb.active; ins.title = "Leia-me"; ins.sheet_view.showGridLines = False
ins.column_dimensions["B"].width = 24; ins.column_dimensions["C"].width = 100
ins["B1"] = "Medição preliminar – Moradia bifamiliar Apúlia (Vendeiro)"; ins["B1"].font = Font(name=F, bold=True, size=15, color="1F3864")
ins["B2"] = "v1 · Elaborada a partir das peças escritas dos projetos (sem MQ do projetista) · Para validação"; ins["B2"].font = Font(name=F, italic=True, size=10, color="595959")
r = 4
for k, (c, t) in CONF.items():
    ins.cell(r, 2, f"Confiança {k}").fill = PatternFill("solid", fgColor=c); ins.cell(r, 2).font = Font(name=F, bold=True)
    ins.cell(r, 3, t).font = Font(name=F); r += 1
r += 1
for t in ["Rastreabilidade: cada quantidade do MQ remete para uma linha da folha «Cálculos» (fórmula + fonte). Altere os dados de base (D01–D20) e tudo recalcula.",
          "Arquitetura: os desenhos do PDF usam fontes não extraíveis; as áreas dos compartimentos foram lidas nas plantas incluídas nos projetos de Gás/Residuais.",
          "Estrutura: volumes de pilares e vigas estimados (secções reais, comprimentos/alturas aproximados). Aço por rácio. Validar com medição sobre desenhos.",
          "Elétrico: desenhos não legíveis → pontos de luz/tomadas por rácio (0,25 e 0,35 por m² útil).",
          "Omissões identificadas e já incluídas: válvula antirretorno (cadastro EAmb), ATE do ITED (qtd 0 no projeto), hidrante SCIE, poço de infiltração pluvial.",
          "Fora de medição (confiança D): demolição do existente, arranjos exteriores, guardas, pintura exterior, luminárias, cozinhas."]:
    c = ins.cell(r, 3, "• " + t); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True); ins.row_dimensions[r].height = 30; r += 1

# Cálculos
ca = wb.create_sheet("Cálculos")
ca["A1"] = "CÁLCULOS AUXILIARES (fórmulas rastreáveis)"; ca["A1"].font = Font(name=F, bold=True, size=13, color="1F3864")
h = ["ID", "Descrição", "Expressão", "", "Resultado", "Un", "Fonte", "Conf."]
for i, t in enumerate(h, 1):
    c = ca.cell(3, i, t); c.font = HFont; c.fill = HF; c.border = B
for cid, desc, expr, un, fonte, conf in CALCULOS:
    rr = ref[cid]
    vals = [cid, desc, expr, None, conv(expr), un, fonte, conf]
    for i, v in enumerate(vals, 1):
        c = ca.cell(rr, i, v); c.font = Font(name=F, size=9); c.border = B; c.alignment = Alignment(wrap_text=i in (2, 3, 7), vertical="top")
    ca.cell(rr, 5).number_format = "#,##0.00"; ca.cell(rr, 5).font = Font(name=F, size=9, bold=True)
    ca.cell(rr, 8).fill = PatternFill("solid", fgColor=CONF[conf][0])
    if cid.startswith("D"): ca.cell(rr, 3).font = Font(name=F, size=9, color="0000FF")
for col, w in zip("ABCDEFGH", [6, 48, 40, 1, 11, 5, 42, 6]): ca.column_dimensions[col].width = w
ca.freeze_panes = "C4"

# MQ
mq = wb.create_sheet("MQ_Preliminar", 1)
mq["A1"] = "MAPA DE QUANTIDADES PRELIMINAR"; mq["A1"].font = Font(name=F, bold=True, size=13, color="1F3864")
mq["A2"] = "Obra: Moradia bifamiliar – Rua da Ponte Nova 30, Apúlia (Esposende) · Requerente: M. A. S. R. Vendeiro"; mq["A2"].font = Font(name=F, italic=True, size=9)
h = ["Artº", "Descrição", "Un", "Quant.", "Origem / cálculo", "Confiança"]
for i, t in enumerate(h, 1):
    c = mq.cell(4, i, t); c.font = HFont; c.fill = HF; c.border = B; c.alignment = Alignment(horizontal="center")
r = 5; cnt = Counter()
for row in MQ:
    if len(row) == 2:
        mq.cell(r, 1, row[0]); mq.cell(r, 2, row[1])
        for i in range(1, 7):
            mq.cell(r, i).fill = PatternFill("solid", fgColor="D9E1F2"); mq.cell(r, i).font = Font(name=F, bold=True, size=10); mq.cell(r, i).border = B
    else:
        cod, desc, un, q, orig, conf = row; cnt[conf] += 1
        vals = [cod, desc, un, conv(q), orig, conf]
        for i, v in enumerate(vals, 1):
            c = mq.cell(r, i, v); c.font = Font(name=F, size=9); c.border = B; c.alignment = Alignment(wrap_text=i in (2, 5), vertical="top")
        mq.cell(r, 4).number_format = "#,##0.00"
        mq.cell(r, 6).fill = PatternFill("solid", fgColor=CONF[conf][0]); mq.cell(r, 6).alignment = Alignment(horizontal="center")
    r += 1
for col, w in zip("ABCDEF", [7, 70, 5, 10, 46, 9]): mq.column_dimensions[col].width = w
mq.freeze_panes = "C5"; mq.auto_filter.ref = f"A4:F{r-1}"
rr = ins.max_row + 2   # resumo na folha Leia-me (fora do MQ, para não ser lido como artigo)
ins.cell(rr, 2, "Resumo da confiança").font = Font(name=F, bold=True)
for k in "ABCD":
    rr += 1; ins.cell(rr, 2, f"{k} – {CONF[k][1]}").fill = PatternFill("solid", fgColor=CONF[k][0]); ins.cell(rr, 3, f"{cnt[k]} artigos")
wb.save(out); print(dict(cnt), "artigos:", sum(cnt.values()))
