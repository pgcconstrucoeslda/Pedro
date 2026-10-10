"""Relatório Word de verificação e compatibilização (para enviar ao projetista/dono de obra)."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = "1F4E78"
COR_GRAV = {"Alta": "FFC7CE", "Média": "FFEB9C", "Baixa": "DDEBF7", "Info": "F2F2F2", "OK": "C6EFCE",
            "CRÍTICO": "FFC7CE", "ALERTA": "FFEB9C"}


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _tabela(doc, df: pd.DataFrame, larguras: list[float] | None = None, cor_col: str | None = None, fonte=8):
    if df is None or df.empty:
        doc.add_paragraph("(sem registos)").italic = True
        return
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(c))
        r.bold = True
        r.font.size = Pt(fonte)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, AZUL)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, c in enumerate(df.columns):
            v = row[c]
            txt = "" if v is None or (isinstance(v, float) and pd.isna(v)) else (f"{v:,.2f}".replace(",", " ") if isinstance(v, float) else str(v))
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(txt)
            run.font.size = Pt(fonte)
        if cor_col and row.get(cor_col) in COR_GRAV:
            for c in cells:
                _shade(c, COR_GRAV[row[cor_col]])
    if not larguras:  # distribuição proporcional ao conteúdo, dentro da largura útil (24,5 cm)
        lens = [max([len(str(c))] + [min(len(str(v)), 60) for v in df[c].head(50)]) for c in df.columns]
        tot = sum(max(l, 4) for l in lens)
        larguras = [24.5 * max(l, 4) / tot for l in lens]
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for j, w in enumerate(larguras[:len(df.columns)]):
        t.columns[j].width = Cm(w)
        for row in t.rows:
            row.cells[j].width = Cm(w)
    doc.add_paragraph()


def gerar_docx(prj, comp: pd.DataFrame | None = None, mq_avisos: list[str] | None = None,
               perfis: list[dict] | None = None, params: dict | None = None,
               regras: pd.DataFrame | None = None) -> bytes:
    """Relatório para enviar ao projetista. Só são listados desvios de regras com
    confiança alta/média (as de confiança baixa são pontos de partida internos)."""
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.6))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    arq = prj.principal()
    leg = {}
    for d in prj.desenhos:
        for k, v in d.res.legenda.items():
            if not k.startswith("_"):
                leg.setdefault(k, v)
    issues = prj.issues_df()

    h = doc.add_heading("Relatório de Verificação de Medições e Compatibilização de Projeto", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = pd.DataFrame([
        ("Projeto", leg.get("DESIGNAÇÃODOPROJECTO", leg.get("DESIGNACAO", "-"))),
        ("Requerente", leg.get("REQUERENTE", "-")),
        ("Localização", leg.get("LOCALIZAÇÃO", "-")),
        ("Fase", leg.get("FASE", "-")),
        ("Peças analisadas", f"{len(prj.desenhos)} desenhos" + (f" + {params.get('Mapa de quantidades')}" if params and params.get("Mapa de quantidades") else "")),
        ("Data do relatório", dt.date.today().strftime("%d/%m/%Y")),
    ], columns=["Campo", "Valor"])
    _tabela(doc, meta, [5, 20], fonte=9)

    # 1. resumo
    doc.add_heading("1. Resumo", 1)
    cnt = issues.gravidade.value_counts() if not issues.empty else pd.Series(dtype=int)
    p = doc.add_paragraph()
    p.add_run("Pontos identificados: ").bold = True
    p.add_run(", ".join(f"{k}: {int(cnt.get(k, 0))}" for k in ("Alta", "Média", "Baixa", "Info", "OK")))
    if comp is not None and not comp.empty:
        cc = comp.estado.value_counts()
        p = doc.add_paragraph()
        p.add_run("Mapa de quantidades vs desenho: ").bold = True
        p.add_run(", ".join(f"{k}: {int(cc.get(k, 0))}" for k in ("OK", "ALERTA", "CRÍTICO", "NÃO MEDIDO", "SEM REGRA")))
    for _, r in issues[issues.gravidade == "Alta"].head(8).iterrows():
        doc.add_paragraph(r.descricao, style="List Bullet")

    # 2. peças
    doc.add_heading("2. Peças analisadas", 1)
    tb = prj.tabela()
    if not tb.empty:
        tb = tb[["desenho", "especialidade", "subespecialidade", "titulo", "data", "revisao", "autores", "projetista",
                 "xrefs_em_falta"]].rename(columns={"xrefs_em_falta": "xrefs em falta"})
    _tabela(doc, tb, [5.5, 3, 3.5, 5, 1.8, 1, 1.5, 3, 4], fonte=7)

    # 3. projetistas
    doc.add_heading("3. Projetistas identificados", 1)
    if perfis:
        ids = {d.perfil_id for d in prj.desenhos}
        pf = pd.DataFrame([dict(perfil=p["id"], nome=p["nome"], evidencia=p.get("evidencia_gabinete", ""),
                                especialidades=", ".join(p["especialidades"]), autores=", ".join(p["autores"]),
                                assinatura=", ".join(p["blocos_legenda"][:3] + p["prefixos_layers"][:3]),
                                desenhos=len(p["projetos"])) for p in perfis if p["id"] in ids or not ids])
        _tabela(doc, pf, [1.3, 4, 5, 4, 3, 6, 1.5], fonte=8)
        doc.add_paragraph("A identificação resulta da assinatura dos desenhos (blocos de legenda, prefixos de layers, "
                          "padrão de nomes) e dos metadados dos ficheiros; deve ser confirmada.").italic = True

    # 4. integridade
    doc.add_heading("4. Integridade das peças entregues", 1)
    _tabela(doc, issues[issues.categoria == "Integridade"][["n", "gravidade", "desenho", "descricao", "evidencia", "acao"]],
            [0.8, 1.5, 3, 9, 6, 5], cor_col="gravidade")

    # 5. níveis e pé-direito
    if arq is not None:
        doc.add_heading("5. Níveis e pé-direito (arquitetura)", 1)
        pdf = arq.res.pe_direito
        if pdf is not None and not pdf.empty:
            _tabela(doc, pdf.rename(columns={"altura_pisos": "altura entre pisos", "pe_direito": "pé-direito",
                                             "metodo_pd": "método", "evidencia_pd": "evidência (handles)"}),
                    [1.5, 2.5, 2, 3, 2.5, 8, 5], fonte=8)
        else:
            doc.add_paragraph("Não foram encontradas cotas de nível suficientes; usado o pé-direito por defeito.")

        # 6. vãos
        doc.add_heading("6. Vãos detetados nas plantas", 1)
        v = arq.res.vaos
        if v is not None and not v.empty:
            cols = [c for c in ["id", "folha", "compartimento", "tipo", "largura", "altura", "area", "codigo_mq"] if c in v]
            doc.add_paragraph(f"{len(v)} vãos detetados ({int((v.tipo == 'porta').sum())} portas, "
                              f"{int((v.tipo != 'porta').sum())} janelas/vãos). As áreas são deduzidas às áreas de paredes.")
            _tabela(doc, v[cols], fonte=7)

    # 7-9 compatibilização / desenho / MQ
    for i, (cat, tit) in enumerate([("Compatibilização", "Compatibilização entre especialidades"),
                                    ("Desenho", "Verificações do desenho de arquitetura"),
                                    ("MQ x Desenho", "Mapa de quantidades vs desenho")], start=7):
        doc.add_heading(f"{i}. {tit}", 1)
        _tabela(doc, issues[issues.categoria == cat][["n", "gravidade", "desenho", "descricao", "evidencia", "acao"]],
                [0.8, 1.5, 3.5, 10, 4, 5], cor_col="gravidade")
        if cat == "MQ x Desenho" and comp is not None and not comp.empty:
            dv = comp[comp.estado.isin(["CRÍTICO", "ALERTA"])].copy()
            n_baixa = 0
            if regras is not None and not regras.empty and "confianca" in regras:
                baixa = set(regras[regras.confianca.astype(str).str.lower() == "baixa"].regra)
                eh_baixa = dv.regras.apply(lambda r: bool(r) and all(x.strip() in baixa for x in str(r).split(",")))
                n_baixa = int(eh_baixa.sum())
                dv = dv[~eh_baixa]
            if n_baixa:
                doc.add_paragraph(f"Nota: {n_baixa} artigo(s) com desvio baseados em regras de confiança baixa "
                                  "(estimativas) não são listados; ver relatório Excel.").italic = True
            if not dv.empty:
                doc.add_paragraph("Artigos com desvio entre o MQ e a medição no desenho:").bold = True
                dv["descricao"] = dv.descricao.astype(str).str.slice(0, 80)
                _tabela(doc, dv[["id", "descricao", "unidade", "qtd_mq_comparada", "qtd_dwg", "desvio_pct", "estado", "observacoes"]]
                        .rename(columns={"qtd_mq_comparada": "MQ", "qtd_dwg": "desenho", "desvio_pct": "desvio %"}),
                        [3.5, 7, 1.2, 1.6, 1.8, 1.6, 1.8, 6], cor_col="estado", fonte=7)

    doc.add_heading("10. Verificações formais do mapa de quantidades", 1)
    for a in mq_avisos or ["(sem observações)"]:
        doc.add_paragraph(a, style="List Bullet")

    # 11. pedidos de esclarecimento
    doc.add_heading("11. Pedidos de esclarecimento / ações", 1)
    acoes = issues[(issues.acao.fillna("") != "") & (issues.gravidade != "OK")]
    if acoes.empty:
        doc.add_paragraph("(sem pedidos)")
    for _, r in acoes.iterrows():
        par = doc.add_paragraph(style="List Number")
        par.add_run(f"[{r.gravidade}] ").bold = True
        par.add_run(f"{r.descricao} ")
        par.add_run(f"→ {r.acao}").italic = True

    doc.add_heading("12. Notas metodológicas", 1)
    for t in [
        "As medições são obtidas automaticamente dos desenhos DWG (convertidos para DXF); cada valor é rastreável até à "
        "entidade do desenho através do seu handle (ver relatório Excel).",
        "Pé-direito: lido das cotas verticais dos cortes que arrancam do nível do pavimento de cada piso; na falta destas, "
        "estimado como altura entre pisos menos a espessura de laje.",
        "Vãos: detetados nos layers de portas/janelas junto aos limites dos compartimentos; a altura é obtida do mapa de "
        "vãos do MQ por correspondência de largura, ou por defeito.",
        "Desenhos que dependem de referências externas não fornecidas não podem ser medidos nem compatibilizados.",
        "Todas as medidas devem ser confirmadas em obra.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
