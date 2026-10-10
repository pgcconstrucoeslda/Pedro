"""Relatório Excel rastreável."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .extractor import Resultado

FILL = {
    "OK": "C6EFCE", "ALERTA": "FFEB9C", "CRÍTICO": "FFC7CE", "NÃO MEDIDO": "D9D9D9",
    "SEM REGRA": "F2F2F2", "SEM ITEM MQ": "BDD7EE",
}
HDR = PatternFill("solid", fgColor="1F4E78")


def _sheet(wb, title, df: pd.DataFrame, widths: dict[str, int] | None = None, freeze="A2"):
    ws = wb.create_sheet(title)
    cols = list(df.columns)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for i, col in enumerate(cols, 1):
        w = (widths or {}).get(col)
        if w is None:
            sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200)] if len(df) else [len(str(col))]
            w = min(max(sample) + 2, 50)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    return ws


def gerar_relatorio(comp: pd.DataFrame, det: pd.DataFrame, regras: pd.DataFrame, regras_res: pd.DataFrame,
                    res: Resultado, mq_avisos: list[str], outros_avisos: list[str], params: dict,
                    prj=None, vaos: pd.DataFrame | None = None, perfis: list[dict] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Relatório de medições DWG x Mapa de Quantidades"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Gerado em", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    for k, v in params.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Estado", "Nº itens"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for st, n in comp.estado.value_counts().items():
        ws.append([st, int(n)])
        if st in FILL:
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=FILL[st])
    ws.append([])
    ws.append(["Avisos e incompatibilidades"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for a in res.avisos + mq_avisos + outros_avisos:
        ws.append(["•", a])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120

    # Comparação com fórmulas (rastreável e editável)
    # A:id B:artigo C:local D:descricao E:unidade F:qtd_mq G:grupo H:qtd_mq_comparada I:qtd_dwg
    # J:desvio K:desvio_pct L:estado
    cols = ["id", "artigo", "local", "descricao", "unidade", "qtd_mq", "grupo", "qtd_mq_comparada", "qtd_dwg",
            "desvio", "desvio_pct", "estado", "regras", "observacoes", "ref_mq"]
    c2 = comp[cols].copy()
    c2["descricao"] = c2.descricao.str.slice(0, 160)
    wc = _sheet(wb, "Comparação", c2, widths={"descricao": 60, "local": 30, "observacoes": 50, "id": 22})
    tol_ok, tol_al = params.get("Tolerância OK (%)", 3), params.get("Tolerância ALERTA (%)", 10)
    for r in range(2, len(c2) + 2):
        if wc[f"G{r}"].value in (None, ""):
            wc[f"H{r}"] = f'=IF(F{r}="","",F{r})'
        wc[f"J{r}"] = f'=IF(OR(H{r}="",I{r}=""),"",I{r}-H{r})'
        wc[f"K{r}"] = f'=IF(OR(H{r}="",I{r}="",H{r}=0),"",ROUND((I{r}-H{r})/H{r}*100,1))'
        if wc[f"L{r}"].value in ("OK", "ALERTA", "CRÍTICO"):
            wc[f"L{r}"] = (f'=IF(K{r}="","NÃO MEDIDO",IF(ABS(K{r})<={tol_ok},"OK",'
                           f'IF(ABS(K{r})<={tol_al},"ALERTA","CRÍTICO")))')
        for col in "FHIJ":
            wc[f"{col}{r}"].number_format = "#,##0.00"
    rng = f"A2:{get_column_letter(len(cols))}{len(c2) + 1}"
    for st, color in FILL.items():
        wc.conditional_formatting.add(rng, FormulaRule(formula=[f'$L2="{st}"'], fill=PatternFill("solid", fgColor=color)))

    if prj is not None:
        iss = prj.issues_df()
        wi = _sheet(wb, "Problemas", iss, widths={"descricao": 90, "evidencia": 40, "acao": 50})
        cores = {"Alta": "FFC7CE", "Média": "FFEB9C", "Baixa": "DDEBF7", "OK": "C6EFCE"}
        for r in range(2, len(iss) + 2):
            c = cores.get(wi.cell(r, 3).value)
            if c:
                for col in range(1, len(iss.columns) + 1):
                    wi.cell(r, col).fill = PatternFill("solid", fgColor=c)
        _sheet(wb, "Peças", prj.tabela())
    _sheet(wb, "Detalhe medições", det.round(3), widths={"descricao": 40})
    _sheet(wb, "Regras", regras, widths={"notas": 60, "filtro": 25})
    if not regras_res.empty:
        _sheet(wb, "Resultado regras", regras_res.round(3), widths={"erro": 60})
    _sheet(wb, "Compartimentos", res.compartimentos.drop(columns=["x", "y"], errors="ignore"))
    if res.pe_direito is not None and not res.pe_direito.empty:
        _sheet(wb, "Níveis e pé-direito", res.pe_direito)
    v = vaos if vaos is not None else res.vaos
    if v is not None and not v.empty:
        _sheet(wb, "Vãos", v.drop(columns=["x", "y"], errors="ignore"))
    if res.xrefs:
        _sheet(wb, "Xrefs", pd.DataFrame(res.xrefs))
    if perfis:
        _sheet(wb, "Projetistas", pd.DataFrame([{
            "id": p["id"], "nome": p["nome"], "gabinete": p.get("gabinete", ""), "evidência": p.get("evidencia_gabinete", ""),
            "especialidades": ", ".join(p["especialidades"]), "autores": ", ".join(p["autores"]),
            "blocos legenda": ", ".join(p["blocos_legenda"]), "prefixos layers": ", ".join(p["prefixos_layers"]),
            "layers conhecidos": "; ".join(f"{k}: {', '.join(v)}" for k, v in p["layers_conhecidos"].items()),
            "blocos conhecidos": "; ".join(f"{k}: {', '.join(v)}" for k, v in p["blocos_conhecidos"].items()),
            "nº documentos": len(p["projetos"])} for p in perfis]))
    p = res.paredes
    if not p.empty:
        resumo = (p.assign(esp_classe=(p.espessura / 0.05).round() * 0.05)
                  .groupby(["folha", "layer", "esp_classe"]).agg(n=("handle", "count"), comprimento=("comprimento", "sum"))
                  .reset_index().round(3))
        _sheet(wb, "Paredes (resumo)", resumo)
        _sheet(wb, "Paredes (tramas)", p)
    _sheet(wb, "Folhas", res.folhas_df().round(2))
    _sheet(wb, "Layers", res.layers_df().round(3))
    blk = res.entidades[res.entidades.tipo == "INSERT"].groupby(["folha", "bloco", "layer"]).size().reset_index(name="n")
    _sheet(wb, "Blocos", blk)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
